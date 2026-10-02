"""Excel 工具箱：读表、统计汇总、按条件筛选、生成 Excel、CSV 互转（openpyxl）。

输入一律是文件空间里的 file_id（对话里「［附件：xxx.xlsx · file_id=…］」那一串，或上一个工具返回的
file_id），支持 .xlsx / .xlsm / .csv；老版 .xls 说明另存为 .xlsx。生成的文件存回文件空间并返回
Markdown 下载链接。任何异常都转成人话返回，不抛给对话；行数、列数、分组数都有上限。

另导出流程积木 ``excel_out``「生成 Excel 表格」：把上一步的表格 / 条目 / 分段写成 xlsx。
"""
from __future__ import annotations

import csv
import datetime as dt
import functools
import io
import logging
import re
import unicodedata
from dataclasses import dataclass
from typing import Literal

from langchain_core.tools import tool
from pydantic import BaseModel, Field, field_validator

from jarvis import files
from jarvis.tenancy import TenantScopeError, current_owner_id

log = logging.getLogger("jarvis")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ROWS = 50000          # 读入的数据行上限
MAX_COLS = 200
MAX_GROUPS = 1000
MAX_CREATE_ROWS = 20000
MAX_CREATE_COLS = 100
SHOW_ROWS = 30            # 结果里最多直接列出的行数
MAX_PREVIEW = 100

FILE_ID_HELP = "文件编号：对话里「［附件：文件名 · file_id=XXX］」中的 XXX，或之前工具返回的 file_id"
SHEET_HELP = "工作表名称或序号（1 开始）；留空用第一个工作表。CSV 文件忽略"
SAVE_HELP = "另存为新 Excel 的文件名，如「部门汇总.xlsx」；留空则只在对话里给结果、不生成文件"


class Problem(Exception):
    """可以直接转告用户的问题。"""


def _friendly(fn):
    """工具体的统一兜底：人话问题原样返回，意外异常只给类名级说明，绝不抛给对话。"""
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        try:
            return fn(*args, **kwargs)
        except Problem as exc:
            return str(exc)
        except Exception as exc:
            log.warning("excel tool %s failed: %s", fn.__name__, type(exc).__name__)
            return ("表格处理没成功（内部错误）。不要用相同参数重试；请告诉领导这一步没成，"
                    "可以检查表格格式、换个条件或稍后再试。")
    return wrapper


def _owner() -> str:
    try:
        return current_owner_id()
    except TenantScopeError as exc:
        raise Problem("当前没有登录的账号，没法读写文件") from exc


def _load(ref: str) -> tuple[dict, bytes]:
    owner = _owner()
    try:
        meta = files.resolve(owner, ref)
        return meta, files.read(owner, meta["id"])
    except KeyError:
        raise Problem(f"没找到文件「{ref}」：可能已过期（文件保留 30 天）或编号不对。"
                      "请让领导重新上传，或使用对话里附件标记中的 file_id。") from None


def _save(owner: str, name: str, data: bytes, mime: str, source: str = "tool") -> dict:
    try:
        return files.save(owner, name, data, mime=mime, source=source)
    except files.FileSpaceError as exc:
        raise Problem(str(exc)) from exc


def _stem(meta: dict) -> str:
    return re.sub(r"\.(xlsx|xlsm|xls|csv)$", "", meta["name"], flags=re.IGNORECASE)


def _done(prefix: str, meta: dict) -> str:
    return f"{prefix}：{files.link(meta)}（file_id={meta['id']}，可继续交给其他工具处理）"


# ---------- 值的解析与显示 ----------

_NUMERIC = re.compile(r"^[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?$")
_PLAIN_NUMBER = re.compile(r"^-?(?:0|[1-9]\d{0,14})(?:\.\d+)?$|^-?[1-9]\d{0,2}(?:,\d{3})+(?:\.\d+)?$")


def to_number(value) -> float | None:
    """单元格 → 数值；「1,234.5」「¥120」「30元」「12%」都认，认不出返回 None。"""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "").replace(" ", "")
    text = re.sub(r"^[¥￥$€£]|元$|块$", "", text)
    percent = text.endswith("%")
    if percent:
        text = text[:-1]
    if not _NUMERIC.match(text):
        return None
    number = float(text)
    return number / 100 if percent else number


def display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.2f}".rstrip("0").rstrip(".")
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d") if (value.hour, value.minute, value.second) == (0, 0, 0) \
            else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    return " ".join(str(value).split())


def markdown_table(headers: list[str], rows: list[list], limit: int = SHOW_ROWS) -> str:
    def cell(value) -> str:
        return display(value).replace("|", "／")[:60]
    lines = ["| " + " | ".join(cell(h) for h in headers) + " |", "|" + " --- |" * len(headers)]
    lines += ["| " + " | ".join(cell(v) for v in row) + " |" for row in rows[:limit]]
    if len(rows) > limit:
        lines.append(f"（共 {len(rows)} 行，这里只列前 {limit} 行）")
    return "\n".join(lines)


def _col_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ---------- 读表 ----------

@dataclass
class Table:
    sheet: str
    headers: list[str]
    rows: list[list]
    truncated: bool = False


def _is_csv(meta: dict) -> bool:
    return meta["name"].lower().endswith(".csv") or meta.get("mime") == "text/csv"


def _check_kind(meta: dict, data: bytes) -> None:
    name = meta["name"].lower()
    if name.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise Problem(f"《{meta['name']}》是老版 Excel（.xls），暂不支持：请在 Excel 或 WPS 里「另存为」.xlsx 后再上传")
    if not _is_csv(meta) and not data.startswith(b"PK"):
        raise Problem(f"《{meta['name']}》不是 Excel（.xlsx）或 CSV 文件，表格工具处理不了")


def _workbook(meta: dict, data: bytes):
    from openpyxl import load_workbook
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise Problem(f"《{meta['name']}》打不开：可能已损坏、设置了密码，或不是真正的 .xlsx 文件") from None


def _pick_sheet(names: list[str], sheet: str) -> str:
    wanted = str(sheet or "").strip()
    if not wanted:
        return names[0]
    if wanted in names:
        return wanted
    for name in names:
        if name.strip().casefold() == wanted.casefold():
            return name
    if wanted.isdigit() and 1 <= int(wanted) <= len(names):
        return names[int(wanted) - 1]
    raise Problem(f"没有叫「{wanted}」的工作表。现有工作表：{'、'.join(names)}")


def _decode(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _raw_rows(meta: dict, data: bytes, sheet: str) -> tuple[str, list[str], list[list], bool]:
    """(工作表名, 全部工作表名, 原始行（含表头，去掉全空行）, 是否截断)。"""
    _check_kind(meta, data)
    raw, truncated = [], False
    if _is_csv(meta):
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for row in csv.reader(io.StringIO(text), dialect):
            if any(cell.strip() for cell in row):
                if len(raw) > MAX_ROWS:
                    truncated = True
                    break
                raw.append([cell.strip() for cell in row[:MAX_COLS]])
        name = _stem(meta)[:31] or "Sheet1"
        return name, [name], raw, truncated
    book = _workbook(meta, data)
    try:
        names = book.sheetnames
        if not names:
            raise Problem(f"《{meta['name']}》里没有工作表")
        chosen = _pick_sheet(names, sheet)
        for row in book[chosen].iter_rows(values_only=True):
            if any(v not in (None, "") for v in row):
                if len(raw) > MAX_ROWS:
                    truncated = True
                    break
                raw.append(list(row[:MAX_COLS]))
        return chosen, names, _trim(raw), truncated
    finally:
        book.close()


def _trim(raw: list[list]) -> list[list]:
    """去掉所有行都为空的尾部列（Excel 的格式区常比数据宽）。"""
    width = 0
    for row in raw:
        for index in range(len(row) - 1, width - 1, -1):
            if row[index] not in (None, ""):
                width = index + 1
                break
    return [(row + [None] * width)[:width] for row in raw]


def _headers(row: list) -> list[str]:
    out, seen = [], {}
    for index, value in enumerate(row):
        title = display(value) or f"列{_col_letter(index)}"
        if title in seen:
            seen[title] += 1
            title = f"{title}_{seen[title]}"
        else:
            seen[title] = 1
        out.append(title)
    return out


def read_table(meta: dict, data: bytes, sheet: str = "") -> tuple[Table, list[str]]:
    """首个非空行当表头；返回 (Table, 全部工作表名)。"""
    chosen, names, raw, truncated = _raw_rows(meta, data, sheet)
    if not raw:
        return Table(chosen, [], [], truncated), names
    raw = _trim(raw)
    return Table(chosen, _headers(raw[0]), raw[1:], truncated), names


def _column(table: Table, name: str) -> int:
    wanted = str(name or "").strip()
    if wanted in table.headers:
        return table.headers.index(wanted)
    folded = [h.strip().casefold() for h in table.headers]
    if wanted.casefold() in folded:
        return folded.index(wanted.casefold())
    letter = re.fullmatch(r"(?:列)?([A-Za-z]{1,2})", wanted)
    if letter:
        target = letter.group(1).upper()
        for index in range(len(table.headers)):
            if _col_letter(index) == target:
                return index
    raise Problem(f"没有找到列「{wanted}」。现有列：{'、'.join(table.headers[:40])}")


def _numeric_column(rows: list[list], index: int) -> bool:
    values = [row[index] for row in rows if row[index] not in (None, "")]
    if not values:
        return False
    numeric = sum(1 for v in values if to_number(v) is not None)
    return numeric >= max(1, 0.6 * len(values))


# ---------- 写表 ----------

def _width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _sheet_title(name: str) -> str:
    title = re.sub(r"[\[\]:*?/\\]", "_", str(name or "").strip())[:31].strip("'")
    return title or "Sheet1"


def cell_value(value):
    """写入前整理：像普通数字的字符串转成数值（方便在 Excel 里求和），其余原样。"""
    if isinstance(value, str):
        text = value.strip().replace("**", "")
        if _PLAIN_NUMBER.match(text):
            number = float(text.replace(",", ""))
            return int(number) if number.is_integer() and "." not in text and abs(number) < 1e15 else number
        return text
    return value


def xlsx_bytes(headers: list[str], rows: list[list], sheet: str = "Sheet1") -> bytes:
    """表头加粗、冻结首行、自动筛选、列宽按内容自适应（中文按两个字符宽）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    ws = book.active
    ws.title = _sheet_title(sheet)
    width = max([len(headers)] + [len(r) for r in rows[:1]]) if (headers or rows) else 0
    header_cells = [str(h) if h is not None else "" for h in headers] + [""] * (width - len(headers))
    widths = [_width(h) for h in header_cells]
    ws.append(header_cells)
    for row in rows:
        values = [cell_value(v) for v in list(row)[:width]]
        ws.append(values)
        for index, value in enumerate(values):
            if index < len(widths):
                widths[index] = max(widths[index], _width(display(value)))
    for cells in ws.iter_rows():   # 「=」开头的文字一律按文字存，不当公式执行
        for item in cells:
            if isinstance(item.value, str) and item.value.startswith("="):
                item.data_type = "s"
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF7")
    for item in ws[1]:
        item.font, item.fill = bold, fill
        item.alignment = Alignment(horizontal="center", vertical="center")
    for index, size in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(60, max(8, size + 2))
    if rows and width:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def parse_markdown_table(text: str) -> tuple[list[str], list[list[str]]] | None:
    """找出文本里第一个 Markdown 表格 → (表头, 行)；没有返回 None。"""
    lines = str(text or "").splitlines()
    for start in range(len(lines) - 1):
        if "|" in lines[start] and _TABLE_RULE.match(lines[start + 1]):
            def cells(line: str) -> list[str]:
                line = line.strip()
                if line.startswith("|"):
                    line = line[1:]
                if line.endswith("|"):
                    line = line[:-1]
                return [c.strip().replace("**", "").replace("`", "") for c in line.split("|")]
            headers = cells(lines[start])
            rows = []
            for line in lines[start + 2:]:
                if "|" not in line or not line.strip():
                    break
                row = cells(line)
                rows.append((row + [""] * len(headers))[: len(headers)])
            return headers, rows
    return None


# ---------- 参数 ----------

_STATS = {"sum": "合计", "mean": "平均", "max": "最大", "min": "最小", "count": "计数"}
_STAT_ALIASES = {
    "求和": "sum", "合计": "sum", "总和": "sum", "总计": "sum", "汇总": "sum", "total": "sum",
    "平均": "mean", "平均值": "mean", "均值": "mean", "avg": "mean", "average": "mean",
    "最大": "max", "最大值": "max", "最小": "min", "最小值": "min",
    "计数": "count", "个数": "count", "数量": "count", "条数": "count",
}
_OPS = {
    "=": "eq", "==": "eq", "等于": "eq", "是": "eq", "eq": "eq",
    "!=": "ne", "<>": "ne", "不等于": "ne", "不是": "ne", "ne": "ne",
    ">": "gt", "大于": "gt", "gt": "gt", ">=": "ge", "大于等于": "ge", "不小于": "ge", "ge": "ge",
    "<": "lt", "小于": "lt", "lt": "lt", "<=": "le", "小于等于": "le", "不大于": "le", "le": "le",
    "contains": "contains", "包含": "contains", "含": "contains",
    "not_contains": "not_contains", "不包含": "not_contains",
    "startswith": "startswith", "开头是": "startswith", "以…开头": "startswith",
    "in": "in", "属于": "in", "是其中之一": "in",
    "empty": "empty", "为空": "empty", "空": "empty", "not_empty": "not_empty", "不为空": "not_empty", "非空": "not_empty",
}
_OP_LABELS = {"eq": "=", "ne": "≠", "gt": ">", "ge": "≥", "lt": "<", "le": "≤", "contains": "包含",
              "not_contains": "不包含", "startswith": "开头是", "in": "属于", "empty": "为空", "not_empty": "不为空"}


class ReadArgs(BaseModel):
    file_id: str = Field(description=FILE_ID_HELP)
    sheet: str = Field(default="", description=SHEET_HELP)
    rows: int = Field(default=20, ge=1, le=MAX_PREVIEW, description="预览前几行数据，默认 20")


class SummaryArgs(BaseModel):
    file_id: str = Field(description=FILE_ID_HELP)
    sheet: str = Field(default="", description=SHEET_HELP)
    group_by: str = Field(default="", description="按哪一列分组汇总，如「部门」；留空则对整列做统计")
    columns: list[str] = Field(default_factory=list, max_length=30,
                               description="要统计的列名，如 [\"金额\"]；留空表示所有数值列")
    stats: list[str] = Field(default_factory=lambda: ["sum"], max_length=5,
                             description="统计方式，可多选：sum 合计 / mean 平均 / max 最大 / min 最小 / count 计数；默认 sum")
    save_as: str = Field(default="", max_length=80, description=SAVE_HELP)

    @field_validator("stats")
    @classmethod
    def _stats(cls, value: list[str]) -> list[str]:
        out = []
        for item in value or ["sum"]:
            key = str(item).strip().lower()
            key = _STAT_ALIASES.get(key, key)
            if key not in _STATS:
                raise ValueError(f"不认识的统计方式「{item}」，只能是 sum / mean / max / min / count")
            if key not in out:
                out.append(key)
        return out or ["sum"]


class Condition(BaseModel):
    column: str = Field(description="列名，如「金额」")
    op: str = Field(description="比较方式：= / != / > / >= / < / <= / contains（包含）/ not_contains / "
                                "startswith / in（属于其中之一，value 用逗号分隔）/ empty / not_empty")
    value: str | float | list[str] = Field(default="", description="比较的值，如 1000、\"销售部\"、\"2026-09-01\"；empty / not_empty 不需要")

    @field_validator("op")
    @classmethod
    def _op(cls, value: str) -> str:
        key = _OPS.get(str(value).strip().lower()) or _OPS.get(str(value).strip())
        if not key:
            raise ValueError(f"不认识的比较方式「{value}」")
        return key


class FilterArgs(BaseModel):
    file_id: str = Field(description=FILE_ID_HELP)
    conditions: list[Condition] = Field(min_length=1, max_length=10, description="筛选条件，1–10 条")
    match: Literal["all", "any"] = Field(default="all", description="all：同时满足全部条件；any：满足任一条件")
    sheet: str = Field(default="", description=SHEET_HELP)
    columns: list[str] = Field(default_factory=list, max_length=50, description="结果只保留这些列；留空保留全部列")
    save_as: str = Field(default="", max_length=80, description="结果文件名；留空自动命名为「原名_筛选结果.xlsx」")


class CreateArgs(BaseModel):
    name: str = Field(default="", max_length=80, description="文件名，如「报名表.xlsx」")
    rows: list[list[str | float | None]] = Field(default_factory=list, max_length=MAX_CREATE_ROWS + 1,
                                  description="表格数据：第一行是表头，后面每行一条记录，如 [[\"姓名\",\"金额\"],[\"张三\",120]]")
    markdown: str = Field(default="", max_length=200000, description="或者直接给一个 Markdown 表格（| 表头 | … |），与 rows 二选一")
    sheet_name: str = Field(default="Sheet1", max_length=31, description="工作表名称")


class ConvertArgs(BaseModel):
    file_id: str = Field(description=FILE_ID_HELP)
    name: str = Field(default="", max_length=80, description="新文件名；留空自动沿用原名")


class ToCsvArgs(ConvertArgs):
    sheet: str = Field(default="", description=SHEET_HELP)


# ---------- 统计 ----------

def _stat(values: list[float], key: str) -> float | None:
    if key == "count":
        return float(len(values))
    if not values:
        return None
    if key == "sum":
        return sum(values)
    if key == "mean":
        return sum(values) / len(values)
    return max(values) if key == "max" else min(values)


def _profile(table: Table, indices: list[int]) -> tuple[list[str], list[list]]:
    headers = ["列", "类型", "非空", "合计", "平均", "最大", "最小", "说明"]
    out = []
    for index in indices:
        values = [row[index] for row in table.rows if row[index] not in (None, "")]
        if _numeric_column(table.rows, index):
            nums = [n for n in (to_number(v) for v in values) if n is not None]
            note = f"{len(values) - len(nums)} 个值不是数字，已跳过" if len(nums) < len(values) else ""
            out.append([table.headers[index], "数值", len(values), _stat(nums, "sum"), _stat(nums, "mean"),
                        _stat(nums, "max"), _stat(nums, "min"), note])
        elif values and all(_date(v) is not None for v in values):
            days = sorted(_date(v) for v in values)
            out.append([table.headers[index], "日期", len(values), None, None, None, None,
                        f"最早 {display(days[0])}，最晚 {display(days[-1])}"])
        else:
            counts: dict[str, int] = {}
            for v in values:
                counts[display(v)] = counts.get(display(v), 0) + 1
            top = sorted(counts.items(), key=lambda kv: -kv[1])[:3]
            note = f"{len(counts)} 种取值" + (f"，最多：{'、'.join(f'{k}（{c}）' for k, c in top)}" if top else "")
            out.append([table.headers[index], "文本", len(values), None, None, None, None, note])
    return headers, out


def _grouped(table: Table, key_index: int, value_indices: list[int], stats: list[str]) -> tuple[list[str], list[list]]:
    groups: dict[str, list[list]] = {}
    for row in table.rows:
        key = display(row[key_index]) or "（空）"
        groups.setdefault(key, []).append(row)
        if len(groups) > MAX_GROUPS:
            raise Problem(f"「{table.headers[key_index]}」有超过 {MAX_GROUPS} 种不同的值，不适合分组汇总；换一列分组试试")
    headers = [table.headers[key_index], "行数"]
    for index in value_indices:
        headers += [f"{table.headers[index]}（{_STATS[s]}）" for s in stats]

    def numbers(rows: list[list], index: int) -> list[float]:
        return [n for n in (to_number(r[index]) for r in rows) if n is not None]

    out = []
    for key, rows in groups.items():
        line = [key, len(rows)]
        for index in value_indices:
            nums = numbers(rows, index)
            line += [_stat(nums, s) for s in stats]
        out.append(line)
    if value_indices:
        out.sort(key=lambda r: (r[2] is None, -(r[2] or 0)))
    else:
        out.sort(key=lambda r: -r[1])
    total = ["合计", len(table.rows)]
    for index in value_indices:
        nums = numbers(table.rows, index)
        total += [_stat(nums, s) for s in stats]
    return headers, out + [total]


@tool("excel_read", args_schema=ReadArgs)
@_friendly
def excel_read(file_id: str, sheet: str = "", rows: int = 20) -> str:
    """查看 Excel / CSV 的结构：有哪些工作表、表头（列名）、共多少行，并预览前 N 行。
    何时用：拿到表格先看清列名再统计 / 筛选；领导问「这张表里有什么」。
    示例：excel_read(file_id="AbC123xyz", sheet="", rows=20)"""
    meta, data = _load(file_id)
    table, names = read_table(meta, data, sheet)
    head = f"《{meta['name']}》"
    if len(names) > 1:
        head += f"共 {len(names)} 个工作表：{'、'.join(names)}；当前看的是「{table.sheet}」"
    if not table.headers:
        return head + "：这个工作表是空的。"
    more = "（超过上限，只读了前 5 万行）" if table.truncated else ""
    body = (f"{head}：{len(table.rows)} 行数据{more} × {len(table.headers)} 列。\n"
            f"列名：{'、'.join(table.headers)}\n\n前 {min(rows, len(table.rows))} 行：\n"
            + markdown_table(table.headers, table.rows, limit=rows))
    return body + f"\n\nfile_id={meta['id']}"


@tool("excel_summary", args_schema=SummaryArgs)
@_friendly
def excel_summary(file_id: str, sheet: str = "", group_by: str = "", columns: list[str] | None = None,
                  stats: list[str] | None = None, save_as: str = "") -> str:
    """统计汇总表格：不分组时按列给出计数 / 合计 / 平均 / 最大 / 最小；给 group_by 时按该列分组汇总（如按部门汇总金额），
    结果按第一项统计从大到小排，末行是合计。给 save_as 就把结果另存成新 Excel 并返回下载链接。
    何时用：「按部门汇总金额」「每个类别卖了多少」「这列平均多少」「汇总后生成一个新表格」。
    示例：excel_summary(file_id="AbC123xyz", group_by="部门", columns=["金额"], stats=["sum"], save_as="部门汇总.xlsx")"""
    stats = stats or ["sum"]
    meta, data = _load(file_id)
    table, _names = read_table(meta, data, sheet)
    if not table.rows:
        return f"《{meta['name']}》的「{table.sheet}」没有数据行，没法统计。"
    chosen = [_column(table, c) for c in (columns or [])]
    if group_by:
        key_index = _column(table, group_by)
        values = [i for i in chosen if i != key_index] or [
            i for i in range(len(table.headers)) if i != key_index and _numeric_column(table.rows, i)]
        headers, rows = _grouped(table, key_index, values[:10], stats)
        title = f"《{meta['name']}》按「{table.headers[key_index]}」分组汇总（{len(rows) - 1} 组，{len(table.rows)} 行数据）"
        if not values:
            title += "；没有找到数值列，只统计了行数"
    else:
        headers, rows = _profile(table, chosen or list(range(min(len(table.headers), 30))))
        title = f"《{meta['name']}》「{table.sheet}」各列统计（{len(table.rows)} 行数据）"
    if table.truncated:
        title += "（表格超过 5 万行，只统计了前 5 万行）"
    text = f"{title}：\n\n{markdown_table(headers, rows, limit=SHOW_ROWS + 1)}"
    if save_as:
        name = files.with_extension(save_as, "xlsx", f"{_stem(meta)}_汇总")
        out = _save(_owner(), name, xlsx_bytes(headers, rows, "汇总"), XLSX_MIME)
        text += "\n\n" + _done("已另存为新 Excel", out)
    return text


def _date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip().replace("/", "-").replace(".", "-")
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", text) or re.match(r"^(\d{4})年(\d{1,2})月(\d{1,2})日", text)
    if not match:
        return None
    try:
        return dt.date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def _matches(cell, op: str, value) -> bool:
    text = display(cell)
    if op == "empty":
        return text == ""
    if op == "not_empty":
        return text != ""
    if op in ("contains", "not_contains"):
        found = str(value).strip().casefold() in text.casefold()
        return found if op == "contains" else not found
    if op == "startswith":
        return text.casefold().startswith(str(value).strip().casefold())
    if op == "in":
        options = value if isinstance(value, list) else re.split(r"[,，、;；]", str(value))
        return text.casefold() in {display(o).strip().casefold() for o in options}
    left_num, right_num = to_number(cell), to_number(value)
    if left_num is not None and right_num is not None:
        left, right = left_num, right_num
    else:
        left_date, right_date = _date(cell), _date(value)
        if left_date is not None and right_date is not None:
            left, right = left_date, right_date
        elif op in ("eq", "ne"):
            left, right = text.casefold(), display(value).strip().casefold()
        else:
            return False
    return {"eq": left == right, "ne": left != right, "gt": left > right, "ge": left >= right,
            "lt": left < right, "le": left <= right}[op]


@tool("excel_filter", args_schema=FilterArgs)
@_friendly
def excel_filter(file_id: str, conditions: list, match: str = "all", sheet: str = "",
                 columns: list[str] | None = None, save_as: str = "") -> str:
    """按条件筛选表格里的行，结果另存为新 Excel 并返回下载链接（同时预览前几行）。
    数字按数值比较，日期按日期比较（如「2026-09-01」），其余按文字比较。
    何时用：「把金额大于 1000 的行筛出来」「只要销售部 9 月以后的记录」「找出电话为空的人」。
    示例：excel_filter(file_id="AbC123xyz", conditions=[{"column":"金额","op":">","value":1000},
    {"column":"部门","op":"=","value":"销售部"}], match="all", save_as="大额订单.xlsx")"""
    meta, data = _load(file_id)
    table, _names = read_table(meta, data, sheet)
    if not table.rows:
        return f"《{meta['name']}》的「{table.sheet}」没有数据行，没法筛选。"
    checks = []
    for item in conditions:
        cond = item if isinstance(item, Condition) else Condition(**item)
        checks.append((_column(table, cond.column), cond))
    test = all if match == "all" else any
    kept = [row for row in table.rows if test(_matches(row[i], c.op, c.value) for i, c in checks)]
    label = (" 且 " if match == "all" else " 或 ").join(
        f"{table.headers[i]}{_OP_LABELS[c.op]}{'' if c.op in ('empty', 'not_empty') else display(c.value)}" for i, c in checks)
    if not kept:
        hint = ""
        for index, cond in checks:
            if cond.op == "eq" and not _numeric_column(table.rows, index):
                sample = []
                for row in table.rows:
                    value = display(row[index])
                    if value and value not in sample:
                        sample.append(value)
                    if len(sample) >= 8:
                        break
                hint += f"\n「{table.headers[index]}」列里的值举例：{'、'.join(sample)}"
        return f"《{meta['name']}》里没有符合「{label}」的行（共 {len(table.rows)} 行）。{hint}"
    keep_cols = [_column(table, c) for c in (columns or [])] or list(range(len(table.headers)))
    headers = [table.headers[i] for i in keep_cols]
    rows = [[row[i] for i in keep_cols] for row in kept]
    name = files.with_extension(save_as, "xlsx", f"{_stem(meta)}_筛选结果") if save_as \
        else files.with_extension(f"{_stem(meta)}_筛选结果", "xlsx", "筛选结果")
    out = _save(_owner(), name, xlsx_bytes(headers, rows, table.sheet), XLSX_MIME)
    return (f"按「{label}」筛出 {len(kept)} 行（共 {len(table.rows)} 行）。\n\n"
            f"{markdown_table(headers, rows, limit=10)}\n\n" + _done("已另存为新 Excel", out))


@tool("excel_create", args_schema=CreateArgs)
@_friendly
def excel_create(name: str = "", rows: list | None = None, markdown: str = "", sheet_name: str = "Sheet1") -> str:
    """把表格数据生成 Excel 文件（.xlsx）并返回下载链接：表头加粗、首行冻结、列宽自适应，像数字的值存成数字。
    何时用：「把这个表格做成 Excel」「把刚才的汇总导出成 Excel」「帮我建一个报名表」。
    rows 第一行是表头；也可以直接传 Markdown 表格。
    示例：excel_create(name="报名表.xlsx", rows=[["姓名","部门","金额"],["张三","销售部",1200]])"""
    headers: list = []
    body: list = []
    if rows:
        headers, body = [display(v) for v in rows[0]], [list(r) for r in rows[1:]]
    elif markdown.strip():
        parsed = parse_markdown_table(markdown)
        if parsed is None:
            return "markdown 里没有找到表格：需要「| 表头1 | 表头2 |」加一行「| --- | --- |」的格式。"
        headers, body = parsed
    else:
        return "请给出表格内容：rows（第一行是表头）或 markdown 表格。"
    if not any(headers):
        return "表头是空的：rows 的第一行要写列名。"
    if len(headers) > MAX_CREATE_COLS or len(body) > MAX_CREATE_ROWS:
        return f"表格太大：最多 {MAX_CREATE_COLS} 列、{MAX_CREATE_ROWS} 行。"
    filename = files.with_extension(name, "xlsx", "表格")
    out = _save(_owner(), filename, xlsx_bytes(headers, body, sheet_name), XLSX_MIME)
    return _done(f"已生成 Excel（{len(body)} 行 × {len(headers)} 列）", out)


@tool("csv_to_excel", args_schema=ConvertArgs)
@_friendly
def csv_to_excel(file_id: str, name: str = "") -> str:
    """把 CSV 文件转成 Excel（.xlsx）：自动识别逗号 / 分号 / 制表符分隔和 UTF-8 / GBK 编码，像数字的值存成数字。
    何时用：「把这个 CSV 转成 Excel」「CSV 用 Excel 打开乱码」。
    示例：csv_to_excel(file_id="AbC123xyz")"""
    meta, data = _load(file_id)
    if not _is_csv(meta):
        return f"《{meta['name']}》不是 CSV 文件；Excel 转 CSV 请用 excel_to_csv。"
    table, _names = read_table(meta, data)
    if not table.headers:
        return f"《{meta['name']}》是空的。"
    filename = files.with_extension(name or _stem(meta), "xlsx", "表格")
    out = _save(_owner(), filename, xlsx_bytes(table.headers, table.rows, _stem(meta)), XLSX_MIME)
    more = "（超过 5 万行，只转了前 5 万行）" if table.truncated else ""
    return _done(f"已把《{meta['name']}》转成 Excel（{len(table.rows)} 行 × {len(table.headers)} 列）{more}", out)


@tool("excel_to_csv", args_schema=ToCsvArgs)
@_friendly
def excel_to_csv(file_id: str, sheet: str = "", name: str = "") -> str:
    """把 Excel 的一个工作表导出成 CSV（UTF-8 带 BOM，Excel 和 WPS 打开中文不乱码）。
    何时用：「导出成 CSV」「系统只收 CSV 格式」。
    示例：excel_to_csv(file_id="AbC123xyz", sheet="Sheet1")"""
    meta, data = _load(file_id)
    if _is_csv(meta):
        return f"《{meta['name']}》本来就是 CSV；要转 Excel 请用 csv_to_excel。"
    chosen, _names, raw, truncated = _raw_rows(meta, data, sheet)
    if not raw:
        return f"《{meta['name']}》的「{chosen}」是空的。"
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    for row in raw:
        writer.writerow([display(v) for v in row])
    filename = files.with_extension(name or f"{_stem(meta)}_{chosen}", "csv", "表格")
    out = _save(_owner(), filename, buffer.getvalue().encode("utf-8-sig"), "text/csv")
    more = "（超过 5 万行，只导出了前 5 万行）" if truncated else ""
    return _done(f"已把「{chosen}」导出成 CSV（{len(raw)} 行）{more}", out)


# ---------- 流程积木：生成 Excel 表格 ----------

_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(\S.*)$")
_BULLET = re.compile(r"^\s*(?:[-*+•·]|\d{1,3}[.、)）])\s+(\S.*)$")


def table_from_context(ctx: dict) -> tuple[list[str], list[list]]:
    """上一步的结果 → 表格：Markdown 表格优先，其次条目（待办）、分段，最后按小节 / 列表 / 段落逐行。"""
    text = ctx.get("text") or ""
    parsed = parse_markdown_table(text)
    if parsed and parsed[1]:
        return parsed
    items = [str(x) for x in (ctx.get("items") or []) if str(x).strip()]
    if items:
        return ["序号", "内容"], [[i, item] for i, item in enumerate(items, 1)]
    parts = ctx.get("parts") or []
    if parts:
        return ["序号", "标题", "内容"], [[i, p.get("title", ""), p.get("text", "")] for i, p in enumerate(parts, 1)]
    rows, section, has_section = [], "", False
    for line in text.splitlines():
        if not line.strip() or set(line.strip()) <= set("-*_ "):
            continue
        heading = _HEADING.match(line)
        if heading:
            section, has_section = heading.group(1).replace("**", "").strip(), True
            continue
        bullet = _BULLET.match(line)
        content = (bullet.group(1) if bullet else line).replace("**", "").replace("`", "").strip()
        rows.append([section, content])
    if has_section:
        return ["序号", "小节", "内容"], [[i, s, c] for i, (s, c) in enumerate(rows, 1)]
    return ["序号", "内容"], [[i, c] for i, (_s, c) in enumerate(rows, 1)]


def run_excel_out(job, ctx: dict, options: dict):
    from jarvis.flows.steps import Outcome, StepFailure, clip, preview

    headers, rows = table_from_context(ctx)
    if not rows:
        raise StepFailure("前面没有可以写进表格的内容")
    rows = rows[:MAX_CREATE_ROWS]
    title = clip((options.get("title") or "").strip() or ctx.get("title") or job.flow.get("name") or "表格", 40)
    try:
        meta = files.save(job.user_id, files.with_extension(title, "xlsx", "表格"),
                          xlsx_bytes(headers, rows, title), mime=XLSX_MIME, source="flow")
    except files.FileSpaceError as exc:
        raise StepFailure(str(exc)) from exc
    ctx.setdefault("links", []).append({"label": "Excel 表格", "url": meta["url"]})
    if job.output is None:
        job.output = {"url": meta["url"], "title": meta["name"], "kind": "file"}
    return Outcome(f"生成了 Excel（{len(rows)} 行）", preview(f"{meta['name']} · {meta['url']}"))


def _steps() -> dict:
    from jarvis.flows.steps import ROLE_OUTPUT, StepSpec
    return {"excel_out": StepSpec(
        "excel_out", "生成 Excel 表格", ROLE_OUTPUT, ("items", "text", "parts"), ("links",),
        ({"key": "title", "label": "文件名", "type": "text", "default": "", "max_length": 40},),
        requires=("files",), run=run_excel_out)}


TOOLS = [excel_read, excel_summary, excel_filter, excel_create, csv_to_excel, excel_to_csv]
STEPS = _steps()
