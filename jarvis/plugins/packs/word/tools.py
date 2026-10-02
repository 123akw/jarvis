"""Word 文档：读出 .docx 的正文与表格，按 Markdown 生成 .docx（python-docx）。

输入是文件空间里的 file_id（对话里「［附件：xxx.docx · file_id=…］」那一串）；生成的文档存回文件空间
并返回 Markdown 下载链接。任何异常都转成人话返回，不抛给对话；字数、行数、表格大小都有上限。

另导出流程积木 ``word_out``「生成 Word 文档」：把上一步整理好的内容写成 .docx。
"""
from __future__ import annotations

import functools
import io
import logging
import re

from langchain_core.tools import tool
from pydantic import BaseModel, Field

from jarvis import files
from jarvis.tenancy import TenantScopeError, current_owner_id

log = logging.getLogger("jarvis")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_READ_CHARS = 20000
DEFAULT_READ_CHARS = 8000
MAX_MARKDOWN_CHARS = 100000
MAX_LINES = 5000
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 30
BODY_FONT, HEADING_FONT = "宋体", "黑体"

FILE_ID_HELP = "文件编号：对话里「［附件：文件名 · file_id=XXX］」中的 XXX，或之前工具返回的 file_id"


class Problem(Exception):
    """可以直接转告用户的问题。"""


def _friendly(fn):
    """工具体的统一兜底：人话问题原样返回，意外异常只给类名级说明，绝不抛给对话。"""
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        try:
            return fn(*args, **kwargs)
        except Problem as exc:
            return str(exc)
        except Exception as exc:
            log.warning("word tool %s failed: %s", fn.__name__, type(exc).__name__)
            return ("Word 文档处理没成功（内部错误）。不要用相同参数重试；请告诉领导这一步没成，"
                    "可以换个文件或稍后再试。")
    return wrapper


def _owner() -> str:
    try:
        return current_owner_id()
    except TenantScopeError as exc:
        raise Problem("当前没有登录的账号，没法读写文件") from exc


def _load(ref: str) -> tuple[dict, bytes]:
    owner = _owner()
    try:
        meta = files.resolve(owner, ref)
        return meta, files.read(owner, meta["id"])
    except KeyError:
        raise Problem(f"没找到文件「{ref}」：可能已过期（文件保留 30 天）或编号不对。"
                      "请让领导重新上传，或使用对话里附件标记中的 file_id。") from None


# ---------- 读 ----------

def _open(meta: dict, data: bytes):
    from docx import Document

    name = meta["name"]
    if name.lower().endswith(".doc") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise Problem(f"《{name}》是老版 Word（.doc），暂不支持：请在 Word 或 WPS 里「另存为」.docx 后再上传")
    if not data.startswith(b"PK"):
        raise Problem(f"《{name}》不是 Word（.docx）文件")
    try:
        return Document(io.BytesIO(data))
    except Exception:
        raise Problem(f"《{name}》打不开：可能已损坏、设置了密码，或不是真正的 .docx 文件") from None


def _cell_text(text: str) -> str:
    return " ".join(str(text or "").split()).replace("|", "／")


def _table_markdown(table) -> str:
    rows = []
    for row in table.rows[:MAX_TABLE_ROWS]:
        cells, last = [], None
        for cell in row.cells[:MAX_TABLE_COLS]:
            if cell._tc is last:   # 横向合并的单元格 python-docx 会重复给出
                continue
            last = cell._tc
            cells.append(_cell_text(cell.text))
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    return "\n".join(lines + ["| " + " | ".join(r) + " |" for r in rows[1:]])


def docx_markdown(document) -> str:
    """按正文顺序把段落（标题 / 列表 / 正文）和表格转成 Markdown。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            text = _table_markdown(Table(child, document))
            if text:
                blocks.append(text)
            continue
        if tag != "p":
            continue
        para = Paragraph(child, document)
        text = para.text.strip()
        if not text:
            continue
        style = ""
        try:
            style = para.style.name if para.style is not None else ""
        except Exception:
            pass
        level = re.match(r"(?:Heading|标题)\s*(\d)", style or "")
        if style == "Title":
            blocks.append(f"# {text}")
        elif level:
            blocks.append("#" * min(6, int(level.group(1)) + 1) + f" {text}")
        elif "List Number" in style:
            blocks.append(f"1. {text}")
        elif "List" in style or child.find(".//{*}numPr") is not None:
            blocks.append(f"- {text}")
        else:
            blocks.append(text)
    return "\n\n".join(blocks)


# ---------- 写 ----------

_RULE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+•]\s+(?:\[[ xX]\]\s*)?(.*)$")
_NUMBER = re.compile(r"^(\s*)\d{1,3}[.)、]\s+(.*)$")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|(?<![*\w])\*[^*\s][^*]*\*(?!\*))")


def _east_asia(style, font: str) -> None:
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), font)


def _runs(paragraph, text: str, bold: bool = False) -> None:
    """行内格式：**粗体**、*斜体*、`代码`、[文字](链接) → 文字（链接）。"""
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if (piece.startswith("**") and piece.endswith("**")) or (piece.startswith("__") and piece.endswith("__")):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 1:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece.startswith("[") and "](" in piece:
            label, url = piece[1:-1].split("](", 1)
            run = paragraph.add_run(f"{label}（{url}）" if url.startswith(("http://", "https://")) else label)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(piece)
        if bold:
            run.bold = True


def _table_cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _add_table(document, header: list[str], rows: list[list[str]]) -> None:
    header = header[:MAX_TABLE_COLS]
    rows = [(r + [""] * len(header))[: len(header)] for r in rows[:MAX_TABLE_ROWS]]
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for col, text in enumerate(header):
        _runs(table.rows[0].cells[col].paragraphs[0], text, bold=True)
    for index, row in enumerate(rows, start=1):
        for col, text in enumerate(row):
            _runs(table.rows[index].cells[col].paragraphs[0], text)


def markdown_docx(markdown: str, title: str = "") -> bytes:
    """Markdown 子集 → .docx：# 标题、段落、- / 1. 列表、| 表格 |、> 引用、``` 代码块、**粗体**。"""
    from docx import Document
    from docx.shared import Pt

    document = Document()
    for name, font in (("Normal", BODY_FONT), ("Title", HEADING_FONT), ("Heading 1", HEADING_FONT),
                       ("Heading 2", HEADING_FONT), ("Heading 3", HEADING_FONT), ("Heading 4", HEADING_FONT)):
        _east_asia(document.styles[name], font)
    document.styles["Normal"].font.size = Pt(11)
    lines = str(markdown or "").replace("\r\n", "\n").split("\n")[:MAX_LINES]
    title = " ".join(str(title or "").split())
    if title:
        document.add_heading(title, 0)
        first = next((ln for ln in lines if ln.strip()), "")
        match = _HEADING.match(first)
        if match and match.group(2).replace("**", "").strip() == title:   # 正文第一行就是同名标题：不重复
            lines.remove(first)
        document.core_properties.title = title[:200]
    index, code = 0, None
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            code = [] if code is None else None
            index += 1
            continue
        if code is not None:
            run = document.add_paragraph().add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(9.5)
            index += 1
            continue
        if not stripped or re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            index += 1
            continue
        if "|" in line and index + 1 < len(lines) and _RULE.match(lines[index + 1]):
            header, rows, index = _table_cells(line), [], index + 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_table_cells(lines[index]))
                index += 1
            _add_table(document, header, rows)
            document.add_paragraph()
            continue
        heading = _HEADING.match(line)
        if heading:
            text = heading.group(2).replace("**", "").strip()
            if text:
                document.add_heading(text, min(4, len(heading.group(1))))
            index += 1
            continue
        bullet, number = _BULLET.match(line), _NUMBER.match(line)
        if bullet or number:
            match = bullet or number
            deep = len(match.group(1).replace("\t", "    ")) >= 2
            style = ("List Bullet" if bullet else "List Number") + (" 2" if deep else "")
            _runs(document.add_paragraph(style=style), match.group(2))
        elif stripped.startswith(">"):
            _runs(document.add_paragraph(style="Quote"), stripped.lstrip(">").strip())
        else:
            _runs(document.add_paragraph(), stripped)
        index += 1
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- 工具 ----------

class ReadArgs(BaseModel):
    file_id: str = Field(description=FILE_ID_HELP)
    max_chars: int = Field(default=DEFAULT_READ_CHARS, ge=500, le=MAX_READ_CHARS, description="最多返回多少字，默认 8000")


class CreateArgs(BaseModel):
    markdown: str = Field(min_length=1, max_length=MAX_MARKDOWN_CHARS,
                          description="文档内容（Markdown）：# 标题、段落、- 列表、1. 编号、| 表格 |、**粗体**")
    name: str = Field(default="", max_length=80, description="文件名，如「活动方案.docx」；留空用标题命名")
    title: str = Field(default="", max_length=100, description="文档大标题（放在第一行）；留空则不加")


@tool("word_read", args_schema=ReadArgs)
@_friendly
def word_read(file_id: str, max_chars: int = DEFAULT_READ_CHARS) -> str:
    """读出 Word 文档（.docx）的正文和表格，按原顺序转成 Markdown（标题、列表、表格都保留）。
    何时用：领导要看 / 总结 / 改写一份 Word，而对话里没有它的全文，或需要看清里面的表格时。
    示例：word_read(file_id="AbC123xyz")"""
    meta, data = _load(file_id)
    text = docx_markdown(_open(meta, data))
    if not text.strip():
        return f"《{meta['name']}》里没有读到文字（可能只有图片）。"
    note = ""
    if len(text) > max_chars:
        text, note = text[:max_chars] + "…", f"（全文 {len(text)} 字，只给出前 {max_chars} 字）"
    return f"《{meta['name']}》的内容{note}：\n\n{text}"


@tool("word_create", args_schema=CreateArgs)
@_friendly
def word_create(markdown: str, name: str = "", title: str = "") -> str:
    """把 Markdown 内容生成 Word 文档（.docx）并返回下载链接：支持标题、段落、列表、编号、表格、引用、粗体。
    何时用：「把刚才的总结生成一份 Word」「写一份通知，给我 Word 文件」「把这份方案导出成文档」。
    示例：word_create(title="活动方案", markdown="## 一、目标\\n- 拉新 200 人\\n\\n| 项目 | 预算 |\\n| --- | --- |\\n| 场地 | 3000 |")"""
    data = markdown_docx(markdown, title)
    default = title or next((m.group(2) for m in map(_HEADING.match, markdown.splitlines()) if m), "") or "文档"
    filename = files.with_extension(name or default, "docx", "文档")
    try:
        meta = files.save(_owner(), filename, data, mime=DOCX_MIME, source="tool")
    except files.FileSpaceError as exc:
        return str(exc)
    return f"已生成 Word 文档：{files.link(meta)}（file_id={meta['id']}）"


# ---------- 流程积木：生成 Word 文档 ----------

def context_markdown(ctx: dict) -> str:
    text = (ctx.get("text") or "").strip()
    if text:
        return text
    parts = ctx.get("parts") or []
    if parts:
        return "\n\n".join(f"## {p.get('title', '')}\n\n{p.get('text', '')}".strip() for p in parts)
    items = [str(x) for x in (ctx.get("items") or []) if str(x).strip()]
    return "\n".join(f"- {x}" for x in items)


def run_word_out(job, ctx: dict, options: dict):
    from jarvis.flows.steps import Outcome, StepFailure, clip, preview

    markdown = context_markdown(ctx)
    if not markdown:
        raise StepFailure("前面没有可以写进文档的内容")
    title = clip((options.get("title") or "").strip() or ctx.get("title") or job.flow.get("name") or "文档", 40)
    try:
        meta = files.save(job.user_id, files.with_extension(title, "docx", "文档"),
                          markdown_docx(markdown[:MAX_MARKDOWN_CHARS], title), mime=DOCX_MIME, source="flow")
    except files.FileSpaceError as exc:
        raise StepFailure(str(exc)) from exc
    ctx.setdefault("links", []).append({"label": "Word 文档", "url": meta["url"]})
    if job.output is None:
        job.output = {"url": meta["url"], "title": meta["name"], "kind": "file"}
    return Outcome("生成了 Word 文档", preview(f"{meta['name']} · {meta['url']}"))


def _steps() -> dict:
    from jarvis.flows.steps import ROLE_OUTPUT, StepSpec
    return {"word_out": StepSpec(
        "word_out", "生成 Word 文档", ROLE_OUTPUT, ("text", "parts", "items"), ("links",),
        ({"key": "title", "label": "文档标题", "type": "text", "default": "", "max_length": 40},),
        requires=("files",), run=run_word_out)}


TOOLS = [word_read, word_create]
STEPS = _steps()
