"""文档解析：把上传的 PDF / Word(docx) / Excel(xlsx) / CSV / TXT / Markdown 提取成纯文本。

上传走 JSON+base64（省掉 multipart 依赖）；docx 用标准库 zipfile 解析
word/document.xml（零依赖），PDF 用 pypdf（纯 Python、零传递依赖），
xlsx 用 openpyxl（只读模式，每个工作表取前若干行转成 Markdown 表格）。
解析结果只在内存里走一遭，注入对话消息；需要原文件的办公插件另走文件空间（jarvis/files.py）。
"""
from __future__ import annotations

import csv
import datetime as dt
import html
import io
import re
import zipfile

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm", ".csv")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024   # 10MB
MAX_DOC_CHARS = 8000                  # 超长截断，避免撑爆模型上下文
TABLE_PREVIEW_ROWS = 60               # 表格每个工作表注入对话的行数（统计 / 筛选交给 Excel 工具读全表）
TABLE_MAX_SHEETS = 5
OLD_EXCEL = "暂不支持老版 Excel（.xls）：请在 Excel 或 WPS 里「另存为」.xlsx 后再上传"
OLD_WORD = "暂不支持老版 Word（.doc）：请在 Word 或 WPS 里「另存为」.docx 后再上传"


class DocumentError(Exception):
    """解析失败；消息可直接展示给用户，不含上游细节。"""


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("gb18030", errors="ignore")


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as bundle:
            xml = bundle.read("word/document.xml").decode("utf-8", "ignore")
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise DocumentError("这个 Word 文档无法解析（只支持 .docx）") from exc
    xml = xml.replace("</w:p>", "\n").replace("<w:tab/>", "\t").replace("<w:br/>", "\n")
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("这个 PDF 无法解析（可能是扫描件或已加密）") from exc


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, dt.datetime):
        value = value.strftime("%Y-%m-%d %H:%M") if (value.hour, value.minute) != (0, 0) else value.strftime("%Y-%m-%d")
    return " ".join(str(value).replace("|", "／").split())


def table_markdown(rows: list, total: int | None = None, limit: int = TABLE_PREVIEW_ROWS) -> str:
    """二维表 → Markdown 表格（首行当表头）；超出 limit 行注明总行数。"""
    rows = [[_cell(v) for v in row] for row in rows if any(v not in (None, "") for v in row)]
    if not rows:
        return "（空表）"
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    shown = rows[: limit + 1]
    lines = ["| " + " | ".join(shown[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(row) + " |" for row in shown[1:]]
    data_rows = (total if total is not None else len(rows)) - 1
    if data_rows > len(shown) - 1:
        lines.append(f"（共 {data_rows} 行数据，以上只列前 {len(shown) - 1} 行）")
    return "\n".join(lines)


def _xlsx_text(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentError("服务器还没装 Excel 解析组件（openpyxl），暂时读不了 Excel") from exc
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError("这个 Excel 文件无法解析（只支持 .xlsx，可能已损坏或有密码）") from exc
    try:
        blocks = []
        names = book.sheetnames
        for name in names[:TABLE_MAX_SHEETS]:
            rows, total = [], 0
            for row in book[name].iter_rows(values_only=True):
                if any(v not in (None, "") for v in row):
                    total += 1
                    if len(rows) <= TABLE_PREVIEW_ROWS:
                        rows.append(list(row))
            blocks.append(f"## 工作表：{name}\n\n{table_markdown(rows, total)}")
        if len(names) > TABLE_MAX_SHEETS:
            blocks.append(f"（还有 {len(names) - TABLE_MAX_SHEETS} 个工作表没列出：{'、'.join(names[TABLE_MAX_SHEETS:])}）")
        return "\n\n".join(blocks)
    finally:
        book.close()


def _csv_text(data: bytes) -> str:
    text = _decode_text(data[3:] if data.startswith(b"\xef\xbb\xbf") else data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows, total = [], 0
    for row in csv.reader(io.StringIO(text), dialect):
        if any(cell.strip() for cell in row):
            total += 1
            if len(rows) <= TABLE_PREVIEW_ROWS:
                rows.append(row)
    return table_markdown(rows, total) if rows else ""


def extract_text(filename: str, data: bytes) -> str:
    """按扩展名解析为纯文本；不支持或解析失败抛 DocumentError。"""
    name = str(filename).lower()
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError("文件超过 10MB 上限")
    if name.endswith(".pdf"):
        text = _pdf_text(data)
    elif name.endswith(".docx"):
        text = _docx_text(data)
    elif name.endswith((".txt", ".md")):
        text = _decode_text(data)
    elif name.endswith((".xlsx", ".xlsm")):
        text = _xlsx_text(data)
    elif name.endswith(".csv"):
        text = _csv_text(data)
    elif name.endswith(".xls"):
        raise DocumentError(OLD_EXCEL)
    elif name.endswith(".doc"):
        raise DocumentError(OLD_WORD)
    else:
        raise DocumentError("只支持 PDF、Word（.docx）、Excel（.xlsx）、CSV、TXT 和 Markdown 文件")
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n")).strip()
    if not text:
        raise DocumentError("没有从文档里读到文字（可能是纯图片扫描件）")
    return text
