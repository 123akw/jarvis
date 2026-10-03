"""第十九轮「流程文件直通」：开始节点的文件字段把原文件存进账号的文件空间，``{{start.<key>_file}}`` 渲染成
附件标记交给 Excel / PDF / Word 工具（files.resolve 认得出）；``{{start.<key>}}`` 仍是读出的文字。

覆盖：原文件进文件空间与去重、标记可 resolve、真 Excel 工具在流程里分组汇总一张小表（数字精确）、PDF 合并
（一行里并排两个原文件）、运行记录 / SSE 带原件链接、各种人话错误（接错变量、空间满、只贴文字、字段名撞名、
原文件引用了非文件字段）、读不出文字但要原文件时照样往下走、模板「Excel 报表分析成 Word」端到端（替身模型）、
一句话生成把文件工具改接原文件并过运行前检查。"""
import io
import json
import re

import jarvis.server  # noqa: F401  注册流程路由
import pytest

from jarvis import files
from jarvis.accounts import AccountStore
from jarvis.flows import compose as C
from jarvis.flows import engine, executor
from jarvis.flows import templates as T
from jarvis.flows.graph import GraphError, validate_graph
from jarvis.flows.store import FlowStore
from jarvis.tenancy import TenantStore, tenant_scope


class Deps:
    def __init__(self, compose=None):
        self.prompts: list[str] = []
        self._compose = compose or (lambda prompt: "AI 的回答")

    def compose(self, user_id, prompt):
        self.prompts.append(prompt)
        return self._compose(prompt)

    def flow_deps(self) -> engine.FlowDeps:
        return engine.FlowDeps(tenant_store=TenantStore, compose=self.compose, feishu_ready=lambda uid: False)


@pytest.fixture()
def owner_id():
    accounts = AccountStore(); accounts._ensure_bootstrap()
    return accounts.list_users()[0]["id"]


def _e(source, target, handle=None):
    return {"source": source, "target": target, "sourceHandle": handle}


def _node(node_id, kind, **data):
    return {"id": node_id, "type": kind, "data": data}


def _start(*fields):
    return {"id": "start", "type": "start", "data": {"fields": list(fields)}}


REPORT = {"key": "report", "label": "Excel 报表", "type": "file", "required": True}


def _run(owner_id, graph, deps=None, inputs=None, name="测试流程"):
    graph = validate_graph(graph)
    store = FlowStore()
    saved = store.create_flow(owner_id, name=name, summary="", graph=graph)
    events = []
    with tenant_scope(owner_id):
        result = executor.execute_graph(flow={"id": saved["id"], "name": name, "graph": graph}, user_id=owner_id,
                                        inputs=inputs or {}, deps=(deps or Deps()).flow_deps(), store=store,
                                        emit=events.append)
    result["flow_id"] = saved["id"]
    return events, result


def _by(events, kind):
    return {e["node_id"]: e for e in events if e["type"] == kind}


DEPTS = ("销售部", "技术部", "市场部")


def _sales_rows(n=200):
    """200 行销售明细：AI 读出的文字只看得到前几十行，Excel 工具要把整张表算对。"""
    return [(DEPTS[i % 3], f"员工{i}", (i % 7) * 100 + (0.5 if i % 2 else 0)) for i in range(n)]


def _xlsx(rows) -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.append(["部门", "姓名", "金额"])
    for row in rows:
        sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _pdf(pages: int) -> bytes:
    from pypdf import PdfWriter
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _fmt(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:.2f}".rstrip("0").rstrip(".")


# ---------- 原文件进文件空间、标记可 resolve ----------

def test_start_file_is_kept_and_file_var_renders_resolvable_marker(owner_id):
    graph = {"nodes": [_start(REPORT),
                       _node("tpl", "template", template="原文件：{{start.report_file}}\n正文：{{start.report}}"),
                       _node("end", "end")],
             "edges": [_e("start", "tpl"), _e("tpl", "end")]}
    data = "部门,金额\n销售部,100\n".encode()
    events, result = _run(owner_id, graph, inputs={"report": {"name": "九月.csv", "data": data}})
    assert result["status"] == "ok", result
    text = result["output"]["text"]
    marker = re.search(r"［附件：九月\.csv · file_id=[A-Za-z0-9_-]+］", text).group(0)
    assert "正文：| 部门 | 金额 |" in text and "| 销售部 | 100 |" in text   # {{start.report}} 仍是读出的文字
    meta = files.resolve(owner_id, marker)
    assert meta["name"] == "九月.csv" and meta["source"] == "flow" and files.read(owner_id, meta["id"]) == data
    start_done = _by(events, "node_done")["start"]
    assert start_done["output"]["files"] == [{"name": "九月.csv", "url": f"/api/files/{meta['id']}", "label": "Excel 报表"}]
    run = FlowStore().list_runs(owner_id, result["flow_id"])[0]   # 运行记录里也带原件
    assert run["nodes"][0]["files"] == start_done["output"]["files"]
    # 同一份文件再跑一次：复用文件空间里那一份，不重复占空间
    _run(owner_id, graph, inputs={"report": {"name": "九月.csv", "data": data}})
    assert [m["id"] for m in files.list(owner_id)] == [meta["id"]]
    # 内容变了就另存一份
    _run(owner_id, graph, inputs={"report": {"name": "九月.csv", "data": data + b"x,1\n"}})
    assert len(files.list(owner_id)) == 2


def test_upload_without_file_var_keeps_office_files_only_and_never_blocks(owner_id, monkeypatch):
    graph = {"nodes": [_start(REPORT), _node("end", "end", output="{{start.report}}")], "edges": [_e("start", "end")]}
    events, _ = _run(owner_id, graph, inputs={"report": {"name": "笔记.md", "data": "# 标题\n内容".encode()}})
    assert "files" not in _by(events, "node_done")["start"]["output"] and files.list(owner_id) == []
    monkeypatch.setattr(files, "QUOTA_BYTES", 10)   # 文件空间满了：没用到原文件的流程照跑
    events, result = _run(owner_id, graph, inputs={"report": {"name": "a.csv", "data": "部门,金额\n销售部,100\n".encode()}})
    assert result["status"] == "ok" and "files" not in _by(events, "node_done")["start"]["output"]


# ---------- 真 Excel 工具在流程里跑通分组汇总 ----------

def _excel_graph(file_arg="{{start.report_file}}"):
    return {"nodes": [_start(REPORT, {"key": "group_by", "label": "按哪一列分组", "type": "text", "default": "部门"}),
                      _node("n1", "tool", title="分组汇总", plugin="excel", tool="excel_summary",
                            args={"file_id": file_arg, "group_by": "{{start.group_by}}", "columns": "金额"}),
                      _node("end", "end", output="{{n1.text}}")],
            "edges": [_e("start", "n1"), _e("n1", "end")]}


def test_excel_summary_runs_on_the_uploaded_original_with_exact_numbers(owner_id):
    rows = _sales_rows()
    events, result = _run(owner_id, _excel_graph(), inputs={"report": {"name": "销售明细.xlsx", "data": _xlsx(rows)}})
    assert result["status"] == "ok", result
    text = result["output"]["text"]
    assert "《销售明细.xlsx》按「部门」分组汇总（3 组，200 行数据）" in text
    for dept in DEPTS:
        amounts = [r[2] for r in rows if r[0] == dept]
        assert f"| {dept} | {len(amounts)} | {_fmt(sum(amounts))} |" in text
    assert f"| 合计 | 200 | {_fmt(sum(r[2] for r in rows))} |" in text


def test_file_arg_wired_to_extracted_text_is_caught_before_running(owner_id):
    deps = Deps()
    events, result = _run(owner_id, _excel_graph("{{start.report}}"), deps,
                          inputs={"report": {"name": "销售明细.xlsx", "data": _xlsx(_sales_rows(5))}})
    assert [e["type"] for e in events] == ["run_start", "node_error", "run_done"]   # 运行前就查出来
    assert events[1]["node_id"] == "n1"
    assert events[1]["message"] == "「分组汇总」要的是文件，请用『开始 · Excel 报表（原文件）』这个变量"
    assert files.list(owner_id) == []


def test_file_arg_without_file_id_after_rendering_is_human(owner_id):
    graph = _excel_graph("{{t.text}}")
    graph["nodes"].insert(1, _node("t", "template", template="随便一段话"))
    graph["edges"] = [_e("start", "t"), _e("t", "n1"), _e("n1", "end")]
    events, result = _run(owner_id, graph, inputs={"report": {"name": "a.csv", "data": "部门,金额\n销售部,1\n".encode()}})
    assert result["status"] == "error"
    assert _by(events, "node_error")["n1"]["message"] == "「分组汇总」要的是文件，请用『开始 · Excel 报表（原文件）』这个变量"


def test_upstream_tool_output_with_file_id_feeds_the_next_file_tool(owner_id):
    """上一个工具另存的新表（返回里带 file_id=…）可以直接交给下一个文件工具。"""
    graph = {"nodes": [_start(REPORT),
                       _node("n1", "tool", title="汇总另存", plugin="excel", tool="excel_summary",
                             args={"file_id": "{{start.report_file}}", "group_by": "部门", "save_as": "部门汇总"}),
                       _node("n2", "tool", title="读汇总表", plugin="excel", tool="excel_read", args={"file_id": "{{n1.text}}"}),
                       _node("end", "end", output="{{n2.text}}")],
             "edges": [_e("start", "n1"), _e("n1", "n2"), _e("n2", "end")]}
    _events, result = _run(owner_id, graph, inputs={"report": {"name": "明细.xlsx", "data": _xlsx(_sales_rows(9))}})
    assert result["status"] == "ok", result
    assert result["output"]["text"].startswith("《部门汇总.xlsx》：4 行数据")


# ---------- PDF：一行里并排两个原文件 ----------

def test_pdf_merge_takes_several_original_files_in_one_line(owner_id):
    """空白 PDF 读不出文字：要原文件的流程照样往下走；第三份没传就只合并两份。"""
    graph = {"nodes": [_start({"key": "a", "label": "第一份", "type": "file", "required": True},
                              {"key": "b", "label": "第二份", "type": "file", "required": True},
                              {"key": "c", "label": "第三份", "type": "file"}),
                       _node("m", "tool", title="合并", plugin="pdf", tool="pdf_merge",
                             args={"file_ids": "{{start.a_file}} {{start.b_file}} {{start.c_file}}", "name": "合起来"}),
                       _node("end", "end", output="{{m.text}}")],
             "edges": [_e("start", "m"), _e("m", "end")]}
    events, result = _run(owner_id, graph, inputs={"a": {"name": "封面.pdf", "data": _pdf(1)},
                                                   "b": {"name": "正文.pdf", "data": _pdf(2)}})
    assert result["status"] == "ok", result
    text = result["output"]["text"]
    assert "已按 《封面.pdf》 → 《正文.pdf》 的顺序合并成 1 个 PDF，共 3 页" in text
    merged = next(m for m in files.list(owner_id) if m["name"] == "合起来.pdf")
    assert any(link["url"] == merged["url"] for link in result["output"]["links"])


# ---------- 人话错误 ----------

def test_full_file_space_is_human_when_original_is_needed(owner_id, monkeypatch):
    monkeypatch.setattr(files, "QUOTA_BYTES", 100)
    events, result = _run(owner_id, _excel_graph(), inputs={"report": {"name": "销售明细.xlsx", "data": _xlsx(_sales_rows(5))}})
    message = _by(events, "node_error")["start"]["message"]
    assert result["status"] == "error"
    assert message.startswith("「开始」：「Excel 报表」的原文件存不进文件空间：文件空间已满")
    assert "请先删掉一些不用的文件再试" in message


def test_pasted_text_cannot_stand_in_for_an_original_file(owner_id):
    events, _ = _run(owner_id, _excel_graph(), inputs={"report": "部门,金额\n销售部,1"})
    assert _by(events, "node_error")["start"]["message"] == \
        "「开始」：「Excel 报表」要上传文件：后面的节点要用它的原文件，只贴文字不行"


def test_unreadable_upload_still_fails_when_only_text_is_used(owner_id):
    graph = {"nodes": [_start(REPORT), _node("end", "end", output="{{start.report}}")], "edges": [_e("start", "end")]}
    events, _ = _run(owner_id, graph, inputs={"report": {"name": "空白.pdf", "data": _pdf(1)}})
    assert _by(events, "node_error")["start"]["message"].startswith("「开始」：没有从文档里读到文字")
    assert files.list(owner_id) == []


def test_original_file_label_in_llm_materials(owner_id):
    graph = {"nodes": [_start(REPORT), _node("ai", "llm", prompt="看看 {{start.report_file}}"), _node("end", "end")],
             "edges": [_e("start", "ai"), _e("ai", "end")]}
    deps = Deps()
    _run(owner_id, graph, deps, inputs={"report": {"name": "a.csv", "data": "部门,金额\n销售部,1\n".encode()}})
    assert '<资料 编号="1" 来源="开始 · Excel 报表（原文件）">\n［附件：a.csv · file_id=' in deps.prompts[0]


@pytest.mark.parametrize("fields, ref, message", [
    ([{"key": "report_file", "label": "报表", "type": "file"}], "", "「开始」的「报表」内部名字不能以 _file 结尾"),
    ([{"key": "report", "label": "报表", "type": "text"}], "{{start.report_file}}",
     "「结束」用到了「报表」的原文件，但它已经不是文件输入了"),
    ([{"key": "report", "label": "报表", "type": "file"}], "{{start.other_file}}", "「结束」用到了「开始」里已经删掉的输入"),
])
def test_validate_graph_checks_original_file_vars(fields, ref, message):
    graph = {"nodes": [_start(*fields), _node("end", "end", output=ref)], "edges": [_e("start", "end")]}
    with pytest.raises(GraphError) as err:
        validate_graph(graph)
    assert str(err.value).startswith(message)


def test_validate_graph_accepts_original_file_var_with_longest_key():
    key = "a" * 24
    graph = {"nodes": [_start({"key": key, "label": "资料", "type": "file"}),
                       _node("end", "end", output="{{start.%s_file}}" % key)], "edges": [_e("start", "end")]}
    assert validate_graph(graph)["nodes"][1]["data"]["output"] == "{{start.%s_file}}" % key


# ---------- 模板：Excel 报表分析成 Word（替身模型） ----------

def test_excel_report_template_sums_with_excel_tool_and_writes_word(owner_id):
    template = T.get_template("excel_report")
    n1 = next(n for n in template["graph"]["nodes"] if n["id"] == "n1")
    assert n1["type"] == "tool" and n1["data"]["tool"] == "excel_summary"
    assert n1["data"]["args"]["file_id"] == "{{start.report_file}}"
    rows = _sales_rows()
    deps = Deps(compose=lambda prompt: "结论：技术部最高。")
    events, result = _run(owner_id, template["graph"], deps, name=template["name"],
                          inputs={"report": {"name": "九月销售.xlsx", "data": _xlsx(rows)}})
    assert result["status"] == "ok", result
    summary = _by(events, "node_done")["n1"]["output"]["text"]
    total = sum(r[2] for r in rows)
    assert f"| 合计 | 200 | {_fmt(total)} |" in summary
    assert f"| 合计 | 200 | {_fmt(total)} |" in deps.prompts[0]   # AI 拿到的是工具算好的数字
    text = result["output"]["text"]
    assert text.startswith("# 金额按部门汇总分析（") and "结论：技术部最高。" in text
    word = next(m for m in files.list(owner_id) if m["name"].endswith(".docx"))
    from docx import Document
    document = Document(io.BytesIO(files.read(owner_id, word["id"])))
    cells = [[c.text for c in row.cells] for table in document.tables for row in table.rows]
    assert ["合计", "200", _fmt(total)] in cells   # Word 里的表格和 Excel 工具算的一致


@pytest.mark.parametrize("template_id", ["pdf_merge", "word_polish"])
def test_new_file_templates_wire_original_files(template_id):
    graph = validate_graph(T.get_template(template_id)["graph"])
    tool = next(n for n in graph["nodes"] if n["type"] == "tool")
    arg = tool["data"]["args"].get("file_id") or tool["data"]["args"].get("file_ids")
    assert "_file}}" in arg


def test_word_polish_template_reads_the_original_docx(owner_id):
    from docx import Document
    document = Document()
    document.add_heading("活动方案", level=1)
    document.add_paragraph("我们打算再周末搞个活动。")
    out = io.BytesIO(); document.save(out)
    deps = Deps(compose=lambda prompt: "# 活动方案\n\n我们打算在周末举办一场活动。")
    events, result = _run(owner_id, T.get_template("word_polish")["graph"], deps,
                          inputs={"doc": {"name": "方案.docx", "data": out.getvalue()}})
    assert result["status"] == "ok", result
    assert "《方案.docx》的内容：\n\n## 活动方案\n\n我们打算再周末搞个活动。" in _by(events, "node_done")["n1"]["output"]["text"]
    assert any(m["name"].startswith("润色后的文档") for m in files.list(owner_id))


# ---------- 一句话生成：文件工具改接原文件，过校验与运行前检查 ----------

class FakeModel:
    def __init__(self, reply):
        self.reply, self.prompts = reply, []

    def __call__(self, user_id, prompt):
        self.prompts.append(prompt)
        return json.dumps(self.reply, ensure_ascii=False)


def _compose(owner_id, reply, description="上传销售表按部门汇总写成 Word"):
    model = FakeModel(reply)
    result = C.compose_draft(owner_id, description, deps=engine.FlowDeps(tenant_store=TenantStore, compose=model))
    return result, model


def test_compose_rewires_file_tool_to_original_file_and_passes_preflight(owner_id):
    reply = {"name": "销售表分析", "fields": [{"key": "sheet", "label": "销售表", "type": "paragraph", "required": True}],
             "nodes": [{"id": "n1", "type": "tool", "title": "汇总", "plugin": "excel", "tool": "excel_summary",
                        "args": {"file_id": "{{start.sheet}}", "group_by": "部门"}},
                       {"id": "n2", "type": "llm", "title": "写分析", "prompt": "分析 {{n1.text}}，原表 {{start.sheet}}"},
                       {"id": "n3", "type": "step", "title": "生成 Word", "step": "word_out"},
                       {"id": "end", "type": "end", "output": "{{n2.text}}"}],
             "edges": [["start", "n1"], ["n1", "n2"], ["n2", "n3"], ["n3", "end"]]}
    result, model = _compose(owner_id, reply)
    assert result["source"] == "model", result["notes"]
    assert len(model.prompts) == 1   # 修补后一次就过了运行前检查，不用再问模型
    graph = validate_graph(result["draft"]["graph"])
    by_id = {n["id"]: n for n in graph["nodes"]}
    assert by_id["start"]["data"]["fields"][0]["type"] == "file"   # 接了文件工具的输入改成文件
    assert by_id["n1"]["data"]["args"]["file_id"] == "{{start.sheet_file}}"
    assert by_id["n2"]["data"]["prompt"] == "分析 {{n1.text}}，原表 {{start.sheet}}"   # AI 那里仍用读出的文字
    assert C.check_runnable(graph, owner_id) == ""


def test_compose_handles_file_suffixed_keys_and_undeclared_original_files(owner_id):
    reply = {"name": "合并 PDF",
             "fields": [{"key": "pdf_a_file", "label": "第一份", "type": "file", "required": True}],
             "nodes": [{"id": "m", "type": "tool", "title": "合并", "plugin": "pdf", "tool": "pdf_merge",
                        "args": {"file_ids": "{{start.pdf_a_file}} {{start.pdf_b_file}}"}},
                       {"id": "end", "type": "end", "output": "{{m.text}}"}],
             "edges": [["start", "m"], ["m", "end"]]}
    result, _ = _compose(owner_id, reply, "把两份 PDF 合成一份")
    assert result["source"] == "model", result["notes"]
    graph = validate_graph(result["draft"]["graph"])
    fields = graph["nodes"][0]["data"]["fields"]
    assert [(f["key"], f["type"]) for f in fields] == [("pdf_a", "file"), ("pdf_b", "file")]
    assert graph["nodes"][1]["data"]["args"]["file_ids"] == "{{start.pdf_a_file}} {{start.pdf_b_file}}"
    assert C.check_runnable(graph, owner_id) == ""
