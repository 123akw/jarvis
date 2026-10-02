"""内置办公插件包：PDF 工具箱 / Excel 工具箱 / Word 文档，以及「生成 Excel 表格」「生成 Word 文档」两块积木。

夹具全部在测试里现场生成（pypdf / openpyxl / python-docx），不依赖仓库里的二进制文件。
"""
import csv
import datetime as dt
import importlib
import io
import json
import re
from pathlib import Path
from types import SimpleNamespace

import pytest

from jarvis import files
from jarvis.flows.steps import ROLE_OUTPUT, StepSpec
from jarvis.tenancy import tenant_scope

pdf = importlib.import_module("jarvis.plugins.packs.pdf.tools")
excel = importlib.import_module("jarvis.plugins.packs.excel.tools")
word = importlib.import_module("jarvis.plugins.packs.word.tools")

PACKS = Path(__file__).resolve().parents[1] / "jarvis" / "plugins" / "packs"
OWNER = "owner-office"
LINK = re.compile(r"\[下载 ([^\]]+)\]\(/api/files/([A-Za-z0-9_-]+)\)")


@pytest.fixture(autouse=True)
def scope():
    with tenant_scope(OWNER):
        yield


def _call(tool_item, **kwargs) -> str:
    result = tool_item.invoke(kwargs)
    assert isinstance(result, str)
    return result


def _links(text: str) -> list[tuple[str, str]]:
    return LINK.findall(text)


def _upload(name: str, data: bytes) -> str:
    return files.save(OWNER, name, data, source="upload")["id"]


# ---------- 夹具 ----------

def make_pdf(pages: list[str], password: str | None = None) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    font = DictionaryObject({NameObject("/F1"): DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica")})})
    for text in pages:
        page = writer.add_blank_page(width=595, height=842)
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 14 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): font})
    if password:
        writer.encrypt(password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def pdf_pages(file_id: str) -> list[str]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(files.read(OWNER, file_id)))
    return [(page.extract_text() or "").strip() for page in reader.pages]


def make_xlsx() -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "明细"
    sheet.append(["日期", "部门", "金额", "备注"])
    for day, dept, amount, note in [
        ("2026-09-01", "销售部", 1200, "差旅"), ("2026-09-03", "市场部", 800, ""),
        ("2026-09-05", "销售部", "1,500", "招待"), ("2026-09-08", "研发部", 300.5, "设备"),
        ("2026-09-12", "市场部", 2000, "投放"),
    ]:
        sheet.append([dt.datetime.fromisoformat(day), dept, amount, note])
    other = book.create_sheet("备注")
    other.append(["说明"])
    other.append(["九月报销"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def load_xlsx(file_id: str):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(files.read(OWNER, file_id)))


def make_docx() -> bytes:
    from docx import Document
    document = Document()
    document.add_heading("项目周报", 1)
    document.add_paragraph("本周完成了接口联调。")
    document.add_paragraph("修复登录问题", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "事项", "负责人"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "上线", "小王"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------- 插件清单（契约第 1 节）----------

@pytest.mark.parametrize("pack", ["pdf", "excel", "word"])
def test_plugin_manifest_matches_entry_module(pack):
    manifest = json.loads((PACKS / pack / "plugin.json").read_text("utf-8"))
    module = importlib.import_module(f"jarvis.plugins.packs.{pack}.tools")
    assert manifest["id"] == pack and re.fullmatch(r"[a-z][a-z0-9_]{1,30}", manifest["id"])
    assert manifest["kind"] == "tool" and manifest["category"] == "documents"
    assert manifest["entry"] == "tools.py" and "files" in manifest["requires"]
    assert manifest["source"] == {"type": "builtin"} and manifest["tier"] == "free"
    assert manifest["tools"] == [t.name for t in module.TOOLS]
    assert manifest["steps"] == list(module.STEPS)
    for name in manifest["tools"]:   # 工具名以插件 id 为前缀（csv_to_excel 是 Excel 工具箱的约定俗成名）
        assert name.startswith(pack) or name == "csv_to_excel"
    for item in module.TOOLS:
        assert "何时用" in item.description and "示例" in item.description
    assert 2 <= len(manifest["examples"]) <= 4 and all("领导" not in e for e in manifest["examples"])


def test_steps_follow_stepspec_contract():
    for module, step_id, name in ((excel, "excel_out", "生成 Excel 表格"), (word, "word_out", "生成 Word 文档")):
        spec = module.STEPS[step_id]
        assert isinstance(spec, StepSpec) and spec.id == step_id and spec.name == name
        assert spec.role == ROLE_OUTPUT and "links" in spec.produces and callable(spec.run)
        assert spec.meta()["options"][0]["key"] == "title"


# ---------- PDF ----------

def test_pdf_info_and_extract_text_by_range():
    file_id = _upload("报告.pdf", make_pdf(["Page one", "Page two", "Page three"]))
    info = _call(pdf.pdf_info, file_id=file_id)
    assert "共 3 页" in info and "A4" in info and "有文字层" in info
    text = _call(pdf.pdf_extract_text, file_id=file_id, pages="2-3")
    assert "第 2 页" in text and "Page two" in text and "Page three" in text and "Page one" not in text
    assert "Page one" in _call(pdf.pdf_extract_text, file_id=file_id)


def test_pdf_merge_keeps_given_order():
    first = _upload("甲.pdf", make_pdf(["A1", "A2"]))
    second = _upload("乙.pdf", make_pdf(["B1"]))
    result = _call(pdf.pdf_merge, file_ids=[second, first], name="合并后")
    (name, new_id), = _links(result)
    assert name == "合并后.pdf" and "共 3 页" in result
    assert pdf_pages(new_id) == ["B1", "A1", "A2"]
    assert files.get(OWNER, new_id)["source"] == "tool"


def test_pdf_split_by_ranges_and_every():
    file_id = _upload("手册.pdf", make_pdf([f"P{i}" for i in range(1, 6)]))
    result = _call(pdf.pdf_split, file_id=file_id, ranges=["1-2", "3-"])
    parts = _links(result)
    assert [n for n, _ in parts] == ["手册_第1-2页.pdf", "手册_第3-5页.pdf"]
    assert pdf_pages(parts[1][1]) == ["P3", "P4", "P5"]
    every = _links(_call(pdf.pdf_split, file_id=file_id, every=2))
    assert [pdf_pages(i) for _, i in every] == [["P1", "P2"], ["P3", "P4"], ["P5"]]


def test_pdf_extract_pages_and_rotate():
    file_id = _upload("合同.pdf", make_pdf(["C1", "C2", "C3", "C4"]))
    (name, picked), = _links(_call(pdf.pdf_extract_pages, file_id=file_id, pages="4,第2页"))
    assert pdf_pages(picked) == ["C4", "C2"] and name == "合同_选取2页.pdf"
    (_, last), = _links(_call(pdf.pdf_extract_pages, file_id=file_id, pages="最后一页"))
    assert pdf_pages(last) == ["C4"]
    result = _call(pdf.pdf_rotate, file_id=file_id, degrees=-90, pages="2")
    (_, rotated), = _links(result)
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(files.read(OWNER, rotated)))
    assert [page.rotation for page in reader.pages] == [0, 270, 0, 0]
    assert "旋转 270 度" in result
    assert "只能是 90" in _call(pdf.pdf_rotate, file_id=file_id, degrees=45)


def test_pdf_problems_are_explained_in_plain_words():
    locked = _upload("加密.pdf", make_pdf(["secret"], password="pw123"))
    assert "密码" in _call(pdf.pdf_info, file_id=locked)
    assert "密码" in _call(pdf.pdf_merge, file_ids=[locked, locked])
    not_pdf = _upload("表.xlsx", make_xlsx())
    assert "不是 PDF" in _call(pdf.pdf_info, file_id=not_pdf)
    broken = _upload("坏.pdf", b"%PDF-1.4 garbage")
    assert "损坏" in _call(pdf.pdf_info, file_id=broken)
    three = _upload("三页.pdf", make_pdf(["x", "y", "z"]))
    assert "超出范围" in _call(pdf.pdf_extract_pages, file_id=three, pages="2-9")
    assert "看不懂" in _call(pdf.pdf_extract_pages, file_id=three, pages="第二页")
    assert "没找到文件" in _call(pdf.pdf_info, file_id="NoSuchFile123")
    many = _upload("多页.pdf", make_pdf([str(i) for i in range(25)]))
    assert "超过一次 20 个" in _call(pdf.pdf_split, file_id=many)


def test_pdf_limits(monkeypatch):
    monkeypatch.setattr(pdf, "MAX_PAGES", 2)
    file_id = _upload("长.pdf", make_pdf(["1", "2", "3"]))
    assert "上限" in _call(pdf.pdf_info, file_id=file_id)
    monkeypatch.setattr(pdf, "MAX_PAGES", 500)
    monkeypatch.setattr(pdf, "MAX_TEXT_PAGES", 1)
    assert "一次最多提取 1 页" in _call(pdf.pdf_extract_text, file_id=file_id)


def test_pdf_scanned_pages_say_so():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    file_id = _upload("扫描.pdf", buffer.getvalue())
    assert "扫描件" in _call(pdf.pdf_extract_text, file_id=file_id)
    assert "可能是扫描件" in _call(pdf.pdf_info, file_id=file_id)


# ---------- Excel ----------

def test_excel_read_lists_sheets_headers_and_preview():
    file_id = _upload("报销.xlsx", make_xlsx())
    text = _call(excel.excel_read, file_id=file_id, rows=2)
    assert "共 2 个工作表：明细、备注" in text and "5 行数据 × 4 列" in text
    assert "列名：日期、部门、金额、备注" in text and "| 2026-09-01 | 销售部 | 1200 | 差旅 |" in text
    assert "只列前 2 行" in text
    assert "九月报销" in _call(excel.excel_read, file_id=file_id, sheet="2")
    assert "没有叫「汇总」的工作表" in _call(excel.excel_read, file_id=file_id, sheet="汇总")


def test_excel_summary_groups_and_saves_new_workbook():
    file_id = _upload("报销.xlsx", make_xlsx())
    text = _call(excel.excel_summary, file_id=file_id, group_by="部门", columns=["金额"],
                 stats=["求和", "mean"], save_as="部门汇总")
    assert "| 市场部 | 2 | 2800 | 1400 |" in text           # 按合计从大到小
    assert "| 销售部 | 2 | 2700 | 1350 |" in text           # 「1,500」这种文本数字也算
    assert "| 合计 | 5 | 5800.5 |" in text
    (name, new_id), = _links(text)
    assert name == "部门汇总.xlsx"
    sheet = load_xlsx(new_id).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("部门", "行数", "金额（合计）", "金额（平均）")
    assert rows[1] == ("市场部", 2, 2800, 1400) and rows[-1][0] == "合计"
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"


def test_excel_summary_profiles_columns_without_group():
    file_id = _upload("报销.xlsx", make_xlsx())
    text = _call(excel.excel_summary, file_id=file_id)
    assert "| 金额 | 数值 | 5 | 5800.5 | 1160.1 | 2000 | 300.5 |" in text
    assert "| 日期 | 日期 | 5 |" in text and "最早 2026-09-01，最晚 2026-09-12" in text
    assert "3 种取值" in text
    assert "没有找到列「金钱」" in _call(excel.excel_summary, file_id=file_id, group_by="金钱")
    assert "不认识的统计方式" in str(pytest.raises(Exception, excel.SummaryArgs, file_id="x", stats=["中位数"]).value)


def test_excel_filter_numbers_dates_text_and_saves():
    file_id = _upload("报销.xlsx", make_xlsx())
    text = _call(excel.excel_filter, file_id=file_id, conditions=[
        {"column": "金额", "op": ">=", "value": 1000}, {"column": "日期", "op": "大于", "value": "2026-09-02"}])
    assert "筛出 2 行" in text and "销售部" in text and "投放" in text
    (name, new_id), = _links(text)
    assert name == "报销_筛选结果.xlsx"
    rows = list(load_xlsx(new_id).active.iter_rows(values_only=True))
    assert rows[0] == ("日期", "部门", "金额", "备注") and len(rows) == 3
    any_text = _call(excel.excel_filter, file_id=file_id, match="any", columns=["部门"], save_as="名单.xlsx",
                     conditions=[{"column": "部门", "op": "in", "value": "研发部,市场部"},
                                 {"column": "备注", "op": "empty"}])
    (_, any_id), = _links(any_text)
    assert [r[0] for r in load_xlsx(any_id).active.iter_rows(values_only=True)] == ["部门", "市场部", "研发部", "市场部"]
    none = _call(excel.excel_filter, file_id=file_id, conditions=[{"column": "部门", "op": "=", "value": "销售"}])
    assert "没有符合" in none and "销售部、市场部、研发部" in none and not _links(none)


def test_excel_create_from_rows_and_markdown():
    text = _call(excel.excel_create, name="报名表", rows=[["姓名", "金额", "编号"], ["张三", "1,200", "007"],
                                                        ["=HYPERLINK(\"x\")", 30.5, None]])
    (name, new_id), = _links(text)
    assert name == "报名表.xlsx" and "2 行 × 3 列" in text
    sheet = load_xlsx(new_id).active
    assert sheet["B2"].value == 1200 and sheet["C2"].value == "007"           # 像数字的存数字，编号保留前导 0
    assert sheet["A3"].data_type == "s" and sheet["A3"].value.startswith("=")  # 「=」开头按文字存，不当公式
    assert sheet["A1"].font.bold and sheet.column_dimensions["A"].width >= 8
    md = _call(excel.excel_create, markdown="说明\n\n| 品类 | **销量** |\n| :-- | --: |\n| 奶茶 | 12 |\n| 咖啡 | 8 |")
    (_, md_id), = _links(md)
    assert list(load_xlsx(md_id).active.iter_rows(values_only=True)) == [("品类", "销量"), ("奶茶", 12), ("咖啡", 8)]
    assert "没有找到表格" in _call(excel.excel_create, markdown="没有表格")
    assert "rows" in _call(excel.excel_create)


def test_csv_and_excel_conversions():
    gbk = "姓名;城市\n张三;上海\n李四;北京\n".encode("gb18030")
    csv_id = _upload("名单.csv", gbk)
    (name, xlsx_id), = _links(_call(excel.csv_to_excel, file_id=csv_id))
    assert name == "名单.xlsx"
    assert list(load_xlsx(xlsx_id).active.iter_rows(values_only=True)) == [("姓名", "城市"), ("张三", "上海"), ("李四", "北京")]
    (csv_name, back_id), = _links(_call(excel.excel_to_csv, file_id=_upload("报销.xlsx", make_xlsx()), sheet="明细"))
    raw = files.read(OWNER, back_id)
    assert csv_name == "报销_明细.csv" and raw.startswith(b"\xef\xbb\xbf")   # UTF-8 BOM：Excel 打开中文不乱码
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0] == ["日期", "部门", "金额", "备注"] and rows[1][:3] == ["2026-09-01", "销售部", "1200"]
    assert "本来就是 CSV" in _call(excel.excel_to_csv, file_id=csv_id)


def test_excel_problems_are_explained_in_plain_words(monkeypatch):
    old = _upload("老表.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    assert "另存为" in _call(excel.excel_read, file_id=old)
    assert "不是 Excel" in _call(excel.excel_read, file_id=_upload("a.pdf", b"%PDF-1.4"))
    assert "打不开" in _call(excel.excel_read, file_id=_upload("坏.xlsx", b"PK\x03\x04broken"))

    def boom(*_args, **_kwargs):
        raise RuntimeError("upstream secret detail")
    monkeypatch.setattr(excel, "read_table", boom)
    text = _call(excel.excel_read, file_id=_upload("报销.xlsx", make_xlsx()))
    assert "内部错误" in text and "secret" not in text


def test_excel_quota_full_is_plain_words(monkeypatch):
    file_id = _upload("报销.xlsx", make_xlsx())
    monkeypatch.setattr(files, "QUOTA_BYTES", files.usage(OWNER) + 10)
    assert "已满" in _call(excel.excel_summary, file_id=file_id, group_by="部门", save_as="x.xlsx")


def test_tools_without_tenant_scope_answer_politely():
    token = None
    from jarvis import tenancy
    token = tenancy._OWNER.set(None)
    try:
        assert "没有登录" in _call(excel.excel_read, file_id="whatever123")
        assert "没有登录" in _call(pdf.pdf_info, file_id="whatever123")
        assert "没有登录" in _call(word.word_create, markdown="# x")
    finally:
        tenancy._OWNER.reset(token)


# ---------- Word ----------

def test_word_read_keeps_headings_lists_and_tables():
    file_id = _upload("周报.docx", make_docx())
    text = _call(word.word_read, file_id=file_id)
    assert "## 项目周报" in text and "本周完成了接口联调。" in text and "- 修复登录问题" in text
    assert "| 事项 | 负责人 |" in text and "| 上线 | 小王 |" in text
    assert "全文" in _call(word.word_read, file_id=file_id, max_chars=500) or len(text) < 500
    assert "老版 Word" in _call(word.word_read, file_id=_upload("旧.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"))
    assert "不是 Word" in _call(word.word_read, file_id=_upload("x.docx", b"plain text"))


def test_word_create_from_markdown_roundtrip():
    markdown = ("# 活动方案\n\n## 一、目标\n- **拉新** 200 人\n  - 老带新\n1. 第一阶段\n\n"
                "| 项目 | 预算 |\n| --- | --- |\n| 场地 | 3000 |\n\n> 注意安全\n\n```\nprint(1)\n```\n参考 [官网](https://example.com)")
    text = _call(word.word_create, markdown=markdown, title="活动方案")
    (name, new_id), = _links(text)
    assert name == "活动方案.docx"
    from docx import Document
    document = Document(io.BytesIO(files.read(OWNER, new_id)))
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs if p.text]
    assert paragraphs[0] == ("Title", "活动方案")
    assert sum(1 for _, t in paragraphs if t == "活动方案") == 1                 # 同名一级标题不重复
    assert ("Heading 2", "一、目标") in paragraphs and ("List Bullet", "拉新 200 人") in paragraphs
    assert ("List Bullet 2", "老带新") in paragraphs and ("List Number", "第一阶段") in paragraphs
    assert ("Quote", "注意安全") in paragraphs and ("Normal", "print(1)") in paragraphs
    assert ("Normal", "参考 官网（https://example.com）") in paragraphs
    bullet = next(p for p in document.paragraphs if p.text == "拉新 200 人")
    assert bullet.runs[0].bold and bullet.runs[0].text == "拉新"
    table = document.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["场地", "3000"] and table.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert document.core_properties.title == "活动方案"
    (auto, _), = _links(_call(word.word_create, markdown="## 会议纪要\n内容"))
    assert auto == "会议纪要.docx"


# ---------- 积木 ----------

def _job(output=None):
    return SimpleNamespace(user_id=OWNER, run_id="run-1", flow={"name": "周报流程"}, payload={},
                           deps=None, store=None, output=output)


def _ctx(**kwargs):
    ctx = {"text": "", "parts": [], "items": [], "title": "", "links": []}
    ctx.update(kwargs)
    return ctx


def test_excel_out_step_writes_table_and_link():
    job, ctx = _job(), _ctx(text="## 本周进展\n- 完成登录\n- 修复导出\n## 下周计划\n- 上线", title="周报")
    outcome = excel.STEPS["excel_out"].run(job, ctx, {"title": ""})
    assert outcome.summary == "生成了 Excel（3 行）"
    link = ctx["links"][-1]
    assert link["label"] == "Excel 表格" and link["url"].startswith("/api/files/")
    assert job.output == {"url": link["url"], "title": "周报.xlsx", "kind": "file"}
    file_id = link["url"].rsplit("/", 1)[1]
    assert files.get(OWNER, file_id)["source"] == "flow"
    rows = list(load_xlsx(file_id).active.iter_rows(values_only=True))
    assert rows == [("序号", "小节", "内容"), (1, "本周进展", "完成登录"), (2, "本周进展", "修复导出"), (3, "下周计划", "上线")]
    table_ctx = _ctx(text="| 品类 | 销量 |\n| --- | --- |\n| 奶茶 | 12 |")
    excel.STEPS["excel_out"].run(_job(), table_ctx, {"title": "销量"})
    assert list(load_xlsx(table_ctx["links"][-1]["url"].rsplit("/", 1)[1]).active.iter_rows(values_only=True)) == [
        ("品类", "销量"), ("奶茶", 12)]
    items_ctx = _ctx(items=["订会议室", "发通知"])
    page = {"url": "/r/token", "title": "结果"}
    job = _job(output=page)
    excel.STEPS["excel_out"].run(job, items_ctx, {})
    assert job.output is page                                            # 已有结果网页就不抢
    from jarvis.flows.steps import StepFailure
    with pytest.raises(StepFailure):
        excel.STEPS["excel_out"].run(_job(), _ctx(), {})


def test_word_out_step_writes_document():
    job, ctx = _job(), _ctx(parts=[{"title": "第一章", "text": "开头"}, {"title": "第二章", "text": "结尾"}])
    outcome = word.STEPS["word_out"].run(job, ctx, {"title": "读书笔记"})
    assert outcome.summary == "生成了 Word 文档" and job.output["title"] == "读书笔记.docx"
    from docx import Document
    document = Document(io.BytesIO(files.read(OWNER, ctx["links"][-1]["url"].rsplit("/", 1)[1])))
    texts = [p.text for p in document.paragraphs if p.text]
    assert texts[:3] == ["读书笔记", "第一章", "开头"]
    from jarvis.flows.steps import StepFailure
    with pytest.raises(StepFailure, match="没有可以写进文档"):
        word.STEPS["word_out"].run(_job(), _ctx(), {})


def test_steps_report_full_space_as_step_failure(monkeypatch):
    from jarvis.flows.steps import StepFailure
    monkeypatch.setattr(files, "QUOTA_BYTES", 1)
    with pytest.raises(StepFailure, match="已满"):
        excel.STEPS["excel_out"].run(_job(), _ctx(items=["a"]), {})
    with pytest.raises(StepFailure, match="已满"):
        word.STEPS["word_out"].run(_job(), _ctx(text="正文"), {})
