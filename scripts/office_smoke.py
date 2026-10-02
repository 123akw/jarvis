#!/usr/bin/env python3
"""办公插件本机冒烟（不进 CI）：临时数据目录里用真实默认模型跑一次
「📎 上传一个 Excel → 让智能体按部门汇总金额并生成新 Excel」，确认模型会调用 Excel 工具箱、
回答里给出下载链接，且生成的表格数字正确。

用法：
    .venv/bin/python scripts/office_smoke.py                       # 默认模型（.env）真跑一轮
    .venv/bin/python scripts/office_smoke.py --two-turns           # 先让它看表、再提汇总要求（多一轮模型调用）
    .venv/bin/python scripts/office_smoke.py --env-file /path/.env --save output/round14-C

- 插件包加载器（第十四轮 B）不在这个分支里：脚本只在本进程内把 PDF / Excel / Word 三个插件包的
  TOOLS 临时并进 Agent 工具表，不改正式注册代码；
- 表格是脚本现场用 openpyxl 拼的报销明细；走真实的 /api/upload（存进文件空间、带附件标记）和 /api/chat；
- 模型默认只跑 1 轮对话（内部工具调用若干次）；密钥只从 .env / 环境变量读，绝不打印；
- 全程用临时目录，跑完即删，不碰真实 data/；不连飞书 / 微信。
模型调用了 Excel 工具、回答里有 /api/files/ 下载链接、链接指向的 xlsx 各部门合计正确时退出码 0。
"""
import argparse
import base64
import importlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

ROWS = [
    ("2026-09-01", "销售部", "王磊", "差旅", 1280),
    ("2026-09-02", "市场部", "陈思", "投放", 5600),
    ("2026-09-03", "研发部", "刘洋", "设备", 3299),
    ("2026-09-05", "销售部", "赵敏", "招待", 860.5),
    ("2026-09-08", "行政部", "孙悦", "办公用品", 432),
    ("2026-09-10", "市场部", "周航", "物料", 1750),
    ("2026-09-12", "研发部", "吴迪", "云资源", 2400),
    ("2026-09-15", "销售部", "王磊", "差旅", 2100),
    ("2026-09-18", "行政部", "郑洁", "快递", 96),
    ("2026-09-22", "市场部", "陈思", "活动场地", 4200),
    ("2026-09-25", "研发部", "刘洋", "培训", 1800),
    ("2026-09-28", "销售部", "赵敏", "差旅", 1540),
]
PACKS = ("pdf", "excel", "word")
LINK = re.compile(r"\]\((/api/files/([A-Za-z0-9_-]{8,64}))\)")


def expected_totals() -> dict[str, float]:
    totals: dict[str, float] = {}
    for _day, dept, _who, _kind, amount in ROWS:
        totals[dept] = totals.get(dept, 0) + amount
    return totals


def build_xlsx() -> bytes:
    import datetime as dt
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "九月报销"
    sheet.append(["日期", "部门", "报销人", "类别", "金额"])
    for day, dept, who, kind, amount in ROWS:
        sheet.append([dt.datetime.fromisoformat(day), dept, who, kind, amount])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def has_model_key() -> bool:
    return any(os.getenv(name, "").strip() for name in ("JARVIS_API_KEY", "DEEPSEEK_API_KEY", "OPENAI_API_KEY"))


def office_tools() -> list:
    tools = []
    for pack in PACKS:
        tools += importlib.import_module(f"jarvis.plugins.packs.{pack}.tools").TOOLS
    return tools


def patch_tool_registry() -> None:
    """只在本进程里把办公插件的工具并进 Agent 工具表（正式注册由插件包加载器负责）。"""
    import jarvis.graph as graph
    original = graph.build_tools
    extra = office_tools()

    def build_tools_with_office(*args, **kwargs):
        base = original(*args, **kwargs)
        names = {item.name for item in base}
        return base + [item for item in extra if item.name not in names]

    graph.build_tools = build_tools_with_office


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--env-file", default=str(ROOT / ".env"), help="读模型密钥的 .env（默认项目根）")
    parser.add_argument("--two-turns", action="store_true", help="先发上传消息让它看表，再单独提汇总要求")
    parser.add_argument("--save", default="", help="把回答和生成的 Excel 存到这个目录（看效果 / 截图用）")
    args = parser.parse_args()

    from dotenv import load_dotenv
    if Path(args.env_file).is_file():
        load_dotenv(args.env_file)   # 已有的环境变量优先
    if not has_model_key():
        print("没找到模型密钥（.env 的 JARVIS_API_KEY 等），这个冒烟需要真模型，退出。")
        return 2
    data_dir = tempfile.mkdtemp(prefix="jarvis-office-smoke-")
    os.environ.update(JARVIS_DATA_DIR=data_dir, JARVIS_ADMIN_USERNAME="smoke", JARVIS_ADMIN_PASSWORD="smoke-pass-2026",
                      JARVIS_SESSION_SECRET="office-smoke-session-secret-at-least-256-bits-long!",
                      JARVIS_ENV="development", JARVIS_ALLOW_INSECURE_COOKIE="1")   # TestClient 走 http
    for name in ("FEISHU_APP_ID", "FEISHU_APP_SECRET"):
        os.environ.pop(name, None)   # 冒烟不连飞书
    try:
        return run(args)
    finally:
        shutil.rmtree(data_dir, ignore_errors=True)


def chat(client, message: str, thread_id: str) -> tuple[str, list[dict]]:
    """发一轮对话，返回 (回答全文, 工具事件)。"""
    answer, events = [], []
    with client.stream("POST", "/api/chat", json={"message": message, "thread_id": thread_id}) as response:
        if response.status_code != 200:
            raise RuntimeError(f"对话失败：HTTP {response.status_code}")
        for line in response.iter_lines():
            if not line.startswith("data: "):
                continue
            event = json.loads(line[6:])
            if event["type"] == "token":
                answer.append(event["text"])
            elif event["type"] in ("tool_start", "tool_result"):
                events.append(event)
            elif event["type"] == "error":
                raise RuntimeError(f"对话失败：{event['message']}")
    return "".join(answer), events


def show_tools(events: list[dict]) -> list[str]:
    called = []
    for event in events:
        if event["type"] == "tool_result":
            called.append(event["name"])
            mark = "✓" if event.get("ok") else "✗"
            detail = " ".join(str(event.get("detail", "")).split())[:110]
            print(f"  {mark} {event['name']}  {event.get('ms') or 0:>5} ms  {detail}")
    return called


def run(args) -> int:
    from fastapi.testclient import TestClient
    patch_tool_registry()
    import jarvis.server as server_mod
    from jarvis import files

    server_mod._initialize_runtime()   # 与正式启动一致：按账号解析模型（这里落到 .env 的默认模型）
    print("模型：用 .env 里的默认模型真跑（密钥不打印）；办公工具已临时并入工具表")
    client = TestClient(server_mod.app)
    if client.post("/api/login", json={"username": "smoke", "password": "smoke-pass-2026"}).status_code != 200:
        print("登录失败"); return 1
    client.headers["X-JWS-CSRF"] = client.get("/api/session").json()["csrf_token"]

    upload = client.post("/api/upload", json={"name": "九月报销明细.xlsx",
                                              "content_b64": base64.b64encode(build_xlsx()).decode()}).json()
    if not upload.get("file"):
        print("上传没存进文件空间：", upload.get("error") or upload.get("file_error")); return 1
    marker = upload["file"]["marker"]
    print(f"上传：{upload['name']} → {marker}（注入文字 {upload['chars']} 字）")
    table = f"【表格开始】\n{upload['text']}\n【表格结束】"
    ask = "请按部门汇总金额（每个部门的合计），并把汇总结果生成一个新的 Excel 文件给我。"
    began = time.monotonic()
    if args.two_turns:
        first = (f"请看看这份表格《{upload['name']}》，先简要说明有哪些工作表和列、大概多少行、主要内容；"
                 f"之后我可能让你统计、筛选或整理。\n{marker}\n\n{table}")
        answer, events = chat(client, first, "office-smoke")
        print("第 1 轮工具：")
        show_tools(events)
        answer, events = chat(client, ask, "office-smoke")
    else:
        answer, events = chat(client, f"{ask}\n{marker}\n\n{table}", "office-smoke")
    print(f"工具调用（{time.monotonic() - began:.1f} 秒）：")
    called = show_tools(events)
    print("回答：\n" + "\n".join(f"  │ {line}" for line in answer.strip().splitlines()))

    ok = True
    if not any(name.startswith("excel_") for name in called):
        print("✗ 模型没有调用 Excel 工具"); ok = False
    links = LINK.findall(answer)
    owner = server_mod._accounts.list_users()[0]["id"]
    produced = None
    for _url, file_id in links:
        try:
            meta = files.get(owner, file_id)
        except KeyError:
            print(f"✗ 回答里的链接指向不存在的文件：{file_id}"); ok = False
            continue
        if meta["name"].lower().endswith(".xlsx") and meta["source"] == "tool":
            produced = meta
    if produced is None:
        print("✗ 回答里没有新生成的 Excel 下载链接"); ok = False
    else:
        response = client.get(produced["url"])
        print(f"下载：{produced['url']} → HTTP {response.status_code}，{len(response.content)} 字节，"
              f"Content-Disposition: {response.headers.get('content-disposition', '')[:80]}")
        ok = ok and check_totals(response.content)
        if args.save:
            target = Path(args.save).expanduser()
            target.mkdir(parents=True, exist_ok=True)
            (target / produced["name"]).write_bytes(response.content)
            (target / "office_smoke_answer.md").write_text(answer, encoding="utf-8")
            print(f"已保存：{target / produced['name']}、{target / 'office_smoke_answer.md'}")
    print("结果：" + ("通过" if ok else "没通过"))
    return 0 if ok else 1


def check_totals(data: bytes) -> bool:
    """生成的表里能找到每个部门，且同一行里有一个数等于该部门的合计金额。"""
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(data), data_only=True).active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    good = True
    for dept, total in expected_totals().items():
        row = next((r for r in rows if dept in [str(v).strip() for v in r if v is not None]), None)
        numbers = [float(v) for v in (row or []) if isinstance(v, (int, float))]
        hit = any(abs(n - total) < 0.01 for n in numbers)
        print(f"  {'✓' if hit else '✗'} {dept} 合计 {total:g}" + ("" if hit else f"（表里是 {row}）"))
        good = good and hit
    return good


if __name__ == "__main__":
    sys.exit(main())
